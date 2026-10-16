from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\doubl\Documents\movielens-content-strategy")
OUT = ROOT / "output" / "weekly_presentations"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "1F4E78"
BLUE = "2F75B5"
PALE_BLUE = "EEF5FA"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9D9D9"
DARK = "1F1F1F"
MUTED = "5F6B76"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run(run, size=10.5, bold=False, italic=False, color=DARK, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def setup_document(doc: Document, week: int):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    ppr = title._element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

    for name, size, before, after in (
        ("Heading 1", 15, 4, 5),
        ("Heading 2", 11.5, 5, 2),
        ("Heading 3", 10.5, 4, 2),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run(f"MOVIELENS CONTENT STRATEGY  |  WEEK {week}"), size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Big Data Analytics  |  Weekly presentation brief")
    set_run(r, size=8, color=MUTED)


def add_bullet(doc: Document, text: str, level=0, size=10.2):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    if not p.runs:
        r = p.add_run(text)
    else:
        p.runs[0].text = text
        r = p.runs[0]
    set_run(r, size=size)
    return p


def add_small_label(doc: Document, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(label), size=9.6, bold=True, color=NAVY)
    set_run(p.add_run(text), size=9.6, color=MUTED)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]], widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(value), size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = widths[idx]
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_run(p.add_run(str(value)), size=font_size)
            if widths:
                cell.width = widths[idx]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cover(doc: Document, week: int, subtitle: str, goal: str, deck_map: list[str], source_anchors: list[str]):
    p = doc.add_paragraph(style="Title")
    p.add_run(f"Week {week} Presentation Content")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(subtitle), size=14, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(16)

    doc.add_heading("Presentation goal", level=1)
    p = doc.add_paragraph(goal)
    p.paragraph_format.space_after = Pt(9)

    doc.add_heading("Deck map", level=1)
    rows = [(str(i + 1), title) for i, title in enumerate(deck_map)]
    add_table(doc, ["Slide", "Purpose"], rows, widths=[Inches(0.72), Inches(6.05)], font_size=9.2)

    doc.add_heading("Evidence to keep open during Q and A", level=1)
    for item in source_anchors:
        add_bullet(doc, item, size=9.4)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    set_run(p.add_run("How to use this file  "), size=9.3, bold=True, color=NAVY)
    set_run(p.add_run("Copy the on-slide text into PowerPoint. Use the visual direction to build the graphic, and rehearse from the speaker cue. Update blank actual-effort fields before presenting."), size=9.3, color=MUTED)
    doc.add_page_break()


def add_slide(doc: Document, number: int, title: str, bullets: list[str], visual: str, cue: str, evidence: str | None = None, table=None, page_break=False):
    doc.add_heading(f"Slide {number}  {title}", level=1)
    doc.add_heading("On-slide copy", level=2)
    for bullet in bullets:
        add_bullet(doc, bullet)
    if table:
        headers, rows, widths, size = table
        add_table(doc, headers, rows, widths=widths, font_size=size)
    add_small_label(doc, "Visual direction  ", visual)
    add_small_label(doc, "Speaker cue  ", cue)
    if evidence:
        add_small_label(doc, "Evidence  ", evidence)
    if page_break:
        doc.add_page_break()
    else:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)


def week1() -> tuple[Path, Document]:
    doc = Document()
    setup_document(doc, 1)
    deck = [
        "Weekly status and time plan", "Project-specific SWOT", "Team communication plan",
        "Problem, stakeholders and objective", "Dataset and decision questions", "Method selection",
        "System architecture", "Four-week work breakdown", "Research paper plan and approval request",
    ]
    add_cover(
        doc, 1, "Project framing, dataset approval and delivery plan",
        "Secure approval for the problem, dataset, decision questions, analytical methods, architecture and four-week delivery plan.",
        deck,
        [
            "requirement/Week_1.pdf and requirement/Big Data Analytics_Final_Project_2026_VNUK.pdf",
            "MovieLens_10M_Final_Project_Plan.md and README.md",
            "Dataset README with 10,000,054 ratings, 10,681 movies and 71,567 total users",
        ],
    )
    add_slide(doc, 1, "This week establishes the project contract", [
        "WHAT: define a decision problem, approve the MovieLens 10M dataset, choose methods and assign four weeks of work.",
        "WHEN: complete scope, architecture and paper outline before data engineering begins in Week 2.",
        "WHY: later modelling and dashboard choices need one agreed decision frame and reproducible pipeline.",
        "PLANNED EFFORT: 8-12 team hours. ACTUAL EFFORT: ______ hours. GAP: ______ hours.",
        "ACTION: record any scope change, owner and due date before closing the meeting.",
    ], "Use a horizontal What / When / Why / Time strip. Keep actual effort blank until the team confirms it.", "Open with the decision the class should approve, then explain why the four-week sequence is feasible.", page_break=False)
    add_slide(doc, 2, "SWOT is specific to this project", [
        "Strengths: 10 million ratings, a clear relational model, reproducible SQL and KNIME workflows, and model outputs written back to MySQL.",
        "Weaknesses: historical data ends in 2009; there is no exposure, viewing, revenue or demographic data.",
        "Opportunities: segment-based curation, discovery campaigns, title-level promotion tests and model governance.",
        "Threats: target leakage, slow full-table operations, cold-start exclusions and overclaiming business impact from ratings alone.",
    ], "Build a 2-by-2 SWOT matrix with one short line per item.", "Connect each weakness or threat to a control later in the deck: chronological splits, TRAIN-only features, support thresholds and stated limitations.", page_break=True)

    add_slide(doc, 3, "Communication plan keeps work auditable", [
        "Project owner: scope, weekly status, decisions and final integration.",
        "Data and SQL owner: schema, import, validation queries and dashboard views.",
        "KNIME and modelling owner: preprocessing, K Means, Logistic Regression and writeback.",
        "Dashboard and report owner: Power BI pages, evidence chain, IEEE paper and presentation assembly.",
        "Cadence: one midweek checkpoint, one end-of-week demo and GitHub updates after each approved change.",
    ], "Use a RACI-style lane diagram. Add team names only after roles are agreed.", "State the escalation rule: data or model changes that affect the decision questions must be raised at the next checkpoint.", table=(["Role", "Member assignment", "Primary handoff"], [
        ("Project owner", "________________", "Approved scope and weekly status"),
        ("Data and SQL", "________________", "Validated tables and SQL evidence"),
        ("KNIME and ML", "________________", "Workflows, metrics and writeback"),
        ("Dashboard and paper", "________________", "Decision views and final narrative"),
    ], [Inches(1.55), Inches(1.65), Inches(3.55)], 8.5), page_break=False)
    add_slide(doc, 4, "The project supports content promotion decisions", [
        "Problem: broad popularity lists can hide titles that fit distinct audience segments.",
        "Primary stakeholders: content strategists, campaign planners, analysts and model reviewers.",
        "Objective: combine audience segments and positive-engagement probabilities to identify titles to promote, maintain, reduce or test.",
        "Success means a reproducible path from raw MovieLens files to evidence-backed dashboard decisions.",
    ], "Show a single decision chain: audience signal -> model evidence -> dashboard -> action.", "Avoid claiming revenue lift. The dataset supports preference and engagement analysis, not causal business impact.", page_break=True)

    add_slide(doc, 5, "MovieLens 10M is large enough for the analytical stack", [
        "10,000,054 ratings from 69,878 rating users across 10,681 movies.",
        "95,580 tags and 21,564 movie-genre relationships enrich content context.",
        "Ratings use a 0.5-5.0 scale; the observed average is 3.512 and 50.06% of ratings are at least 4.0.",
        "The time range is 9 January 1995 to 5 January 2009, which supports chronological validation.",
        "Decision questions: Which audience segments exist? Which users are likely to rate positively? Which titles have high predicted appeal but lower historical popularity? Which results are reliable enough to act on?",
    ], "Use four KPI tiles plus the four decision questions beneath them.", "Justify the dataset with scale, relational structure and timestamps; immediately note its historical and non-commercial limits.", evidence="Dataset README and SQL quality-check output.", page_break=False)
    add_slide(doc, 6, "Two models answer different decision questions", [
        "K Means groups warm-start users by activity, rating behavior and genre affinity.",
        "Logistic Regression estimates whether a rating will be positive, defined as rating >= 4.0.",
        "K Means supports audience strategy; Logistic Regression supports title-level prioritization and transparent evaluation.",
        "A chronological TRAIN / VALID / TEST design limits future information from entering historical features.",
    ], "Use a two-column method-to-decision map. Put K Means on the left and Logistic Regression on the right.", "Explain that method choice follows the decision question: segmentation for audience structure, classification for positive engagement.", page_break=True)

    add_slide(doc, 7, "The architecture separates computation from decisions", [
        "MovieLens files -> MySQL raw and curated tables -> KNIME preprocessing and modelling.",
        "KNIME writes segments, predictions and model metrics back to MySQL.",
        "Power BI reads versioned result tables and views for five decision-oriented pages.",
        "GitHub stores SQL, workflow documentation, dashboard assets and the research paper.",
    ], "Use tmp/report/figures/architecture.png as the base visual.", "Walk left to right. Emphasize that MySQL is both the governed input layer and the model-output contract.", evidence="Repository architecture plan, SQL scripts, KNIME assets and dashboard PBIP files.", page_break=False)
    add_slide(doc, 8, "The work breakdown follows four evidence gates", [
        "Week 1: approve scope, dataset, decision questions, architecture and paper outline.",
        "Week 2: validate data, build the MySQL schema, connect KNIME and create the modelling dataset.",
        "Week 3: engineer leakage-safe features, evaluate both models, write results to MySQL and draft the dashboard.",
        "Week 4: finish the dashboard, recommendations, reproducible repository, IEEE paper and Q and A.",
    ], "Use a four-lane timeline with an evidence gate at the end of each week.", "Treat these hours as a planning baseline, not reported effort.", table=(["Week", "Planning baseline", "Exit evidence"], [
        ("1", "8-12 team hours", "Approved project contract"),
        ("2", "16-20 team hours", "Validated SQL-KNIME pipeline"),
        ("3", "20-24 team hours", "Evaluated ML and dashboard draft"),
        ("4", "16-20 team hours", "Final system, paper and presentation"),
    ], [Inches(0.7), Inches(1.7), Inches(4.35)], 8.7), page_break=True)

    add_slide(doc, 9, "The paper documents one reproducible decision system", [
        "Working title: Segment-Based Movie Promotion and Content Discovery Strategy Using MovieLens 10M.",
        "Outline: Abstract, Introduction, Related Work, Methodology, Experiments and Results, Discussion and Recommendations, Conclusion, References.",
        "Initial references: MovieLens dataset documentation; Harper and Konstan on MovieLens; Lloyd on least-squares clustering; official MySQL, KNIME and Power BI documentation.",
        "Approval requested: problem, dataset, four decision questions, K Means plus Logistic Regression, architecture and four-week plan.",
    ], "End with a checklist that the instructor can approve verbally.", "Ask for specific objections to scope or method. Record changes as decisions, owners and dates.", evidence="Final-project brief requires an IEEE-style paper of 6-10 pages with at least five sources.", page_break=False)
    path = OUT / "Week_1_Project_Planning_Slide_Content.docx"
    return path, doc


def week2() -> tuple[Path, Document]:
    doc = Document()
    setup_document(doc, 2)
    deck = [
        "Weekly status and time review", "Dataset and decision questions", "Data quality results",
        "MySQL schema and import", "SQL evidence and KPIs", "KNIME integration and preparation",
        "Initial findings", "Paper progress and Week 3 plan",
    ]
    add_cover(doc, 2, "Data quality, MySQL and KNIME integration",
              "Demonstrate a validated dataset, working MySQL-KNIME connection, reproducible preparation and a modelling-ready chronological split.",
              deck,
              [
                  "python_sql_scripts/results/03_quality_checks_results.xlsx",
                  "sql/ and python_sql_scripts/ for schema, import and quality checks",
                  "knime/ plus screenshots/knime/movielens_like_prediction.png",
              ])
    add_slide(doc, 1, "Week 2 converts raw files into governed analytical data", [
        "WHAT: validate the data, load MySQL, connect KNIME, prepare features and define the chronological split.",
        "WHEN: complete quality evidence and modelling tables before Week 3 experiments begin.",
        "WHY: model metrics are not credible unless the inputs, keys and split logic can be reproduced.",
        "PLANNED EFFORT: 16-20 team hours. ACTUAL EFFORT: ______ hours. GAP: ______ hours.",
        "ACTION FROM GAP: ________________________________________________________________",
    ], "Use a progress bar from raw files to modelling-ready tables.", "Report work completed against the Week 1 baseline. Do not estimate actual hours during the presentation.", page_break=False)
    add_slide(doc, 2, "The dataset supports four decisions", [
        "Audience: Which stable user segments can guide curation?",
        "Engagement: Which user-title pairs are likely to produce ratings of 4.0 or higher?",
        "Portfolio: Which titles combine high predicted appeal with lower historical popularity?",
        "Governance: Which results have enough coverage, support and model quality for action?",
        "Core tables: ratings, movies, tags and a normalized movie_genres bridge.",
    ], "Use a central schema surrounded by four decision questions.", "Keep the business questions visible before discussing technical cleaning steps.", page_break=True)

    add_slide(doc, 3, "Quality checks found no record-level defects", [
        "Missing, orphan and invalid-value checks returned zero affected ratings and zero affected tags.",
        "Duplicate user-movie ratings returned no rows; every rating user had at least 20 ratings.",
        "10,677 of 10,681 catalogue movies have ratings, leaving four catalogue titles without rating activity.",
        "No rows were deleted or imputed. Before and after row counts are unchanged.",
        "Transformations created the genre bridge, chronological split and TRAIN-only aggregate features.",
    ], "Use a before-and-after quality table with green zero-defect indicators and one note for four unrated titles.", "Separate validation from transformation: the data passed integrity checks, while modelling still required leakage-safe restructuring.", evidence="03_quality_checks_results.xlsx and the executed SQL outputs.", table=(["Check", "Result", "Treatment"], [
        ("Ratings", "10,000,054 before and after", "No deletion or imputation"),
        ("Orphan ratings", "0", "No corrective action"),
        ("Orphan tags", "0", "No corrective action"),
        ("Invalid ratings or timestamps", "0", "No corrective action"),
        ("Catalogue coverage", "10,677 of 10,681 rated", "Retain four unrated titles"),
    ], [Inches(2.0), Inches(2.05), Inches(2.7)], 8.5), page_break=False)
    add_slide(doc, 4, "The MySQL design preserves keys and reusable outputs", [
        "ratings uses user_id plus movie_id as the natural interaction key and stores rating and timestamp.",
        "movies stores title and parsed release year; movie_genres resolves the multi-value genre field.",
        "tags stores free-text annotations linked to users and movies.",
        "Indexes support joins and time-based filtering; analytical result tables remain separate from raw data.",
    ], "Show a compact relational diagram: movies 1-to-many ratings, users 1-to-many ratings, movies many-to-many genres.", "Explain why normalized genres improve filtering and aggregation in both KNIME and Power BI.", page_break=True)

    add_slide(doc, 5, "SQL makes the evidence traceable", [
        "SELECT and WHERE validate ranges, timestamps and inclusion rules.",
        "GROUP BY produces rating distributions, activity counts and genre KPIs.",
        "JOIN connects ratings to movies and genres while preserving referential integrity.",
        "Dashboard views expose stable business fields instead of repeating complex joins in Power BI.",
        "Core KPIs: 10.00M ratings, 69,878 rating users, 10,681 movies, average rating 3.512 and 50.06% positive ratings.",
    ], "Use one short SQL excerpt next to five KPI tiles. Highlight the join keys, not the full query.", "Tie every KPI to the query that produced it so the audience can distinguish evidence from interpretation.", page_break=False)
    add_slide(doc, 6, "KNIME reads from and writes to MySQL", [
        "Database Connector and Reader nodes pull governed tables from MySQL.",
        "Preparation nodes enforce types, create the positive-rating target and assemble model fields.",
        "The workflow creates chronological TRAIN, VALID and TEST partitions rather than a random split.",
        "Database Writer nodes persist segments, predictions and metrics for dashboard use.",
        "Credentials remain outside version control; the repository keeps reproducible workflow logic and setup notes.",
    ], "Use a KNIME node-flow screenshot with arrows labeled read, transform, split and writeback.", "Demonstrate the live connection briefly, then return to the documented pipeline rather than navigating every node.", evidence="screenshots/knime/movielens_like_prediction.png and repository KNIME documentation.", page_break=True)

    add_slide(doc, 7, "Initial findings shape the modelling design", [
        "The target is balanced enough for classification: 50.06% of all ratings are positive.",
        "User activity is uneven, so the segmentation model must standardize features before K Means.",
        "The time span allows a realistic out-of-time test rather than evaluating on randomly mixed history.",
        "Tag volume is small relative to ratings, so tags remain contextual evidence rather than a primary model feature.",
        "Cold-start coverage must be reported because aggregate features require prior TRAIN history.",
    ], "Use four finding-to-design arrows: balance, scale, time and coverage.", "Present these as design implications, not model results; evaluation belongs in Week 3.", page_break=False)
    add_slide(doc, 8, "Week 3 starts from a frozen modelling contract", [
        "Paper progress: problem, dataset, architecture and data-quality methodology are ready for the draft.",
        "Week 3 plan: build TRAIN-only features, tune K Means, fit Logistic Regression and evaluate on TEST.",
        "Write user_segments, rating_predictions and model_metrics back to MySQL.",
        "Build an initial five-page Power BI dashboard from versioned result views.",
        "Carry-forward issue or decision: ____________________________________________________",
    ], "Use a checklist with owners and evidence links. Place the planned-versus-actual time line at the bottom.", "Close by naming the exact exit evidence required next week: evaluated models, SQL writeback and dashboard draft.", table=(["Effort item", "Planned", "Actual", "Gap and action"], [
        ("Data and SQL", "______ h", "______ h", "________________________"),
        ("KNIME preparation", "______ h", "______ h", "________________________"),
        ("Documentation", "______ h", "______ h", "________________________"),
    ], [Inches(1.65), Inches(1.05), Inches(1.05), Inches(3.0)], 8.5), page_break=False)
    path = OUT / "Week_2_Data_Quality_MySQL_KNIME_Slide_Content.docx"
    return path, doc


def week3() -> tuple[Path, Document]:
    doc = Document()
    setup_document(doc, 3)
    deck = [
        "Weekly status and time review", "Experimental design and leakage controls", "K Means workflow",
        "K selection and segment profiles", "Logistic Regression workflow", "Classification results",
        "MySQL writeback", "Initial dashboard", "Paper progress and Week 4 plan",
    ]
    add_cover(doc, 3, "Machine learning, evaluation and dashboard draft",
              "Present functioning and evaluated K Means and Logistic Regression models, persist their results in MySQL and show an initial decision dashboard.",
              deck,
              [
                  "kmeans/3.csv through kmeans/7.csv and tmp/report/figures/clustering_summary.png",
                  "Model metrics and confusion matrix in the repository result files",
                  "MySQL tables user_segments, rating_predictions and model_metrics",
              ])
    add_slide(doc, 1, "Week 3 turns governed data into evaluated evidence", [
        "WHAT: engineer features, evaluate K Means and Logistic Regression, write results to MySQL and connect the dashboard.",
        "WHEN: freeze model outputs before Week 4 recommendations and final reporting.",
        "WHY: decisions need out-of-time performance, segment interpretation and a traceable writeback contract.",
        "PLANNED EFFORT: 20-24 team hours. ACTUAL EFFORT: ______ hours. GAP: ______ hours.",
        "ACTION FROM GAP: ________________________________________________________________",
    ], "Use a four-stage strip: features -> models -> evaluation -> dashboard.", "Start with the experimental contract, then show results. This prevents the metrics from appearing detached from methodology.", page_break=False)
    add_slide(doc, 2, "Chronological splits and TRAIN-only features limit leakage", [
        "TRAIN: 7,000,042 ratings from 9 January 1995 to 2 April 2005.",
        "VALID: 1,500,006 ratings from 2 April 2005 to 28 December 2006.",
        "TEST: 1,500,006 ratings from 28 December 2006 to 5 January 2009.",
        "TRAIN rows use leave-one-out user and movie aggregates; VALID and TEST use frozen TRAIN aggregates.",
        "user_id and movie_id are identifiers and are not Logistic Regression predictors.",
    ], "Use a time axis with hard split boundaries and a TRAIN-only feature arrow.", "Explain the rule in one sentence: no row may use information that would not have existed at its prediction time.", page_break=True)

    add_slide(doc, 3, "K Means segments warm-start users", [
        "Create user-level TRAIN features for activity, average rating and genre affinity.",
        "Standardize numeric features so high-volume fields do not dominate Euclidean distance.",
        "Compare K values from 3 through 7 using mean silhouette and cluster interpretability.",
        "Assign business labels only after inspecting final centroids and behavior profiles.",
        "Persist model version, cluster number and final segment label for each eligible user.",
    ], "Show the KNIME pipeline beside a small feature table. Avoid listing every node.", "Clarify that segment names summarize observed profiles; they are not demographic identities.", page_break=False)
    add_slide(doc, 4, "K equals 4 balances separation and interpretability", [
        "Mean silhouette: K=3 0.06021; K=4 0.06864; K=5 0.06919; K=6 0.06460; K=7 0.06511.",
        "K=5 improves mean silhouette by only 0.00055 over K=4, while K=4 gives clearer decision profiles.",
        "54,873 users are segmented, equal to 78.5% of rating users.",
        "Genre Specialists are the largest segment at 17,346 users and have a 71.43% positive-rating rate.",
        "Selective Critics are the most active, averaging 385.3 TRAIN ratings per user.",
    ], "Use tmp/report/figures/clustering_summary.png. Highlight the selected K=4 point and two segment facts.", "A low silhouette means the segments overlap. Use them as decision lenses, not as perfectly separated natural groups.", table=(["Segment", "Users", "Positive rate", "Profile cue"], [
        ("Diverse Explorers", "12,182", "51.38%", "Thriller and action affinity"),
        ("Selective Critics", "10,888", "50.88%", "Highest activity"),
        ("Mainstream Enthusiasts", "14,457", "42.06%", "Family and animation affinity"),
        ("Genre Specialists", "17,346", "71.43%", "Largest and most positive"),
    ], [Inches(2.15), Inches(1.0), Inches(1.25), Inches(2.35)], 8.4), page_break=True)

    add_slide(doc, 5, "Logistic Regression predicts positive engagement", [
        "Target: 1 when rating >= 4.0; otherwise 0.",
        "Predictors use prior user behavior, prior movie behavior, activity and context derived from TRAIN.",
        "TRAIN leave-one-out aggregation prevents each rating from leaking into its own predictors.",
        "VALID selects settings and TEST supplies the final reported metrics.",
        "Prediction rows store probability, class, actual outcome, split and model version.",
    ], "Use a feature-to-probability pipeline ending in a 0.5 decision threshold.", "Logistic Regression was chosen because coefficients and probabilities are easier to explain and govern than a black-box alternative.", page_break=False)
    add_slide(doc, 6, "The model beats the majority-class baseline on warm-start TEST rows", [
        "Accuracy 69.26% versus 59.32% baseline, an improvement of 9.93 percentage points.",
        "AUC 75.83%, precision 60.58%, recall 69.94% and F1 64.92%.",
        "Confusion matrix: TN 31,273; FP 14,189; FN 9,369; TP 21,802.",
        "TEST feature coverage is 76,633 of 1,500,006 rows, or 5.11%; 94.89% are excluded by cold-start rules.",
        "The performance claim therefore applies only to covered warm-start rows.",
    ], "Use tmp/report/figures/model_evaluation.png and place coverage beside the metrics, not in a footnote.", "Lead with baseline improvement, then immediately qualify it with coverage. This is the central governance message.", page_break=True)

    add_slide(doc, 7, "Model outputs return to MySQL with version fields", [
        "user_segments stores user_id, cluster, segment label and model version.",
        "rating_predictions stores probability, predicted class, actual class, split and version.",
        "model_metrics stores metric name, metric value, evaluation split and run metadata.",
        "Dashboard views join stable dimensions to these result tables without recalculating the models.",
        "The writeback step creates one auditable handoff from analytics to reporting.",
    ], "Show three result tables feeding one dashboard-view layer.", "Demonstrate row counts and one sample record from each output table rather than scrolling through raw results.", page_break=False)
    add_slide(doc, 8, "The dashboard draft follows five decisions", [
        "Executive Overview: dataset scale, positive rate, model status and key actions.",
        "Audience Segments: segment size, behavior and genre affinity.",
        "Positive Engagement Model: baseline comparison, metrics, confusion matrix and coverage.",
        "Content Opportunity Matrix: predicted appeal, actual response, popularity and support.",
        "Model Governance: versions, split logic, thresholds, coverage and limitations.",
        "Global controls: model version, segment, genre and minimum support.",
    ], "Use tmp/report/figures/dashboard_map.png as a five-page navigation map.", "Explain what decision each page supports. Avoid describing visuals without linking them to an action.", page_break=True)

    add_slide(doc, 9, "Week 4 will convert model evidence into bounded recommendations", [
        "Paper progress: methodology, experimental design and preliminary results are ready for the IEEE draft.",
        "Finalize Power BI interactions and validate that every visual uses versioned MySQL views.",
        "Create at least three recommendations with evidence, owner, test design, assumptions and risks.",
        "Complete repository instructions, final report, live demo path and Q and A backup slides.",
        "Carry-forward issue or decision: ____________________________________________________",
    ], "Use a Week 4 acceptance checklist and the effort table below.", "End by agreeing on the freeze point for data, model version and dashboard refresh.", table=(["Effort item", "Planned", "Actual", "Gap and action"], [
        ("Feature and model work", "______ h", "______ h", "________________________"),
        ("Writeback and dashboard", "______ h", "______ h", "________________________"),
        ("Paper and rehearsal", "______ h", "______ h", "________________________"),
    ], [Inches(1.75), Inches(1.0), Inches(1.0), Inches(3.0)], 8.5), page_break=False)
    path = OUT / "Week_3_Machine_Learning_Evaluation_Slide_Content.docx"
    return path, doc


def week4() -> tuple[Path, Document]:
    doc = Document()
    setup_document(doc, 4)
    deck = [
        "Final status and time review", "End-to-end system demonstration", "Final dashboard",
        "Audience-segment findings", "Model performance and governance", "Content opportunities",
        "Recommendations for promotion", "Recommendations for testing and controls",
        "Final submission and reproducibility", "Limitations, reflection and Q and A",
    ]
    add_cover(doc, 4, "Final dashboard, recommendations and reproducible submission",
              "Demonstrate the complete system, defend model quality and limits, and present evidence-backed actions that can be tested safely.",
              deck,
              [
                  "dashboard/ for the Power BI project and semantic model",
                  "MySQL writeback tables and dashboard views in sql/",
                  "report/ and output/ for the final paper and deliverables",
              ])
    add_slide(doc, 1, "Week 4 closes the analytical and reporting loop", [
        "WHAT: demonstrate the full pipeline, final dashboard, recommendations, paper and reproducible repository.",
        "WHEN: freeze the model version and dashboard data before the final presentation.",
        "WHY: the final assessment requires one defensible chain from source data to decision.",
        "PLANNED EFFORT: 16-20 team hours. ACTUAL EFFORT: ______ hours. GAP: ______ hours.",
        "ACTION FROM GAP: ________________________________________________________________",
    ], "Use a completion checklist with one evidence link per deliverable.", "State the final model version and data refresh time at the start of the live demo.", page_break=False)
    add_slide(doc, 2, "The system runs from raw files to governed decisions", [
        "MovieLens files are validated and loaded into MySQL.",
        "KNIME reads curated tables, creates leakage-safe features and runs K Means plus Logistic Regression.",
        "Segments, predictions and metrics are written back to versioned MySQL tables.",
        "Power BI reads dashboard views and exposes evidence through filters and five pages.",
        "Recommendations use model evidence only after coverage, support and limitations are checked.",
    ], "Use tmp/report/figures/architecture.png. Animate one stage at a time if presenting in PowerPoint.", "Demonstrate one traceable record or KPI across the layers. The purpose is proof of integration, not a tour of every script.", page_break=True)

    add_slide(doc, 3, "Five dashboard pages separate decisions from governance", [
        "Executive Overview answers whether the project has usable evidence and where action is concentrated.",
        "Audience Segments explains who the model groups and how those groups differ.",
        "Positive Engagement Model shows accuracy, discrimination, errors and coverage.",
        "Content Opportunity Matrix compares predicted appeal, actual response, historical popularity and support.",
        "Model Governance records versions, thresholds, split logic and limitations.",
        "Filters for model version, segment, genre and minimum support keep comparisons consistent.",
    ], "Use tmp/report/figures/dashboard_map.png, then switch to the live PBIP dashboard.", "Show the same filter path across pages so the audience sees how one question becomes a governed decision.", page_break=False)
    add_slide(doc, 4, "Segments support different curation strategies", [
        "Genre Specialists are the largest segment and have the highest positive-rating rate at 71.43%.",
        "Selective Critics provide the strongest title-level support because they average 385.3 TRAIN ratings per user.",
        "Diverse Explorers show relative thriller and action affinity.",
        "Mainstream Enthusiasts show relative family and animation affinity but the lowest positive rate at 42.06%.",
        "K=4 was retained for interpretability even though K=5 had a slightly higher silhouette.",
    ], "Use the segment profile panel from tmp/report/figures/clustering_summary.png.", "Translate profiles into curation hypotheses. Do not treat the labels as fixed identities or causal explanations.", page_break=True)

    add_slide(doc, 5, "The engagement model adds lift but has narrow TEST coverage", [
        "Accuracy is 69.26%, 9.93 percentage points above the 59.32% baseline.",
        "AUC is 75.83%; F1 is 64.92%; precision is 60.58%; recall is 69.94%.",
        "The model evaluates 76,633 TEST rows, equal to 5.11% coverage.",
        "Cold-start exclusions account for 94.89% of TEST rows, so the model is a warm-start decision aid.",
        "Governance controls: chronological splits, TRAIN-only aggregates, minimum support and model-version filters.",
    ], "Use tmp/report/figures/model_evaluation.png with a large 5.11% coverage label.", "Answer performance questions with both quality and applicability: how well the model ranks covered rows, and how many rows it covers.", page_break=False)
    add_slide(doc, 6, "The opportunity matrix distinguishes promotion from overexposure", [
        "Only Selective Critics has title-segment pairs with support >= 50: 107 qualifying pairs.",
        "Hidden opportunities include Hotel Rwanda, My Neighbor Totoro, City of God and Oldboy.",
        "The Incredibles has high predicted appeal but a lower actual positive rate, so it needs testing rather than automatic promotion.",
        "Overexposed candidates include Home Alone, Star Wars Episode I, Twister and The Mask.",
        "Support thresholds prevent unstable small samples from driving the recommendation list.",
    ], "Use tmp/report/figures/content_opportunity.png and label the four action quadrants.", "Explain that opportunity is a prioritization signal. It does not prove that a campaign will cause higher engagement.", table=(["Title", "Support", "Predicted", "Actual", "TRAIN popularity"], [
        ("Hotel Rwanda", "108", "76.72%", "77.78%", "215"),
        ("My Neighbor Totoro", "58", "74.87%", "74.14%", "487"),
        ("City of God", "60", "71.66%", "70.00%", "968"),
        ("Oldboy", "52", "69.94%", "67.31%", "72"),
    ], [Inches(2.1), Inches(0.85), Inches(1.05), Inches(0.9), Inches(1.5)], 8.3), page_break=True)

    add_slide(doc, 7, "Promote validated hidden opportunities to Selective Critics", [
        "Recommendation 1: test Hotel Rwanda and My Neighbor Totoro in Selective Critics placements because predicted and actual positive rates both exceed 74% with support above 50.",
        "Recommendation 2: add City of God and Oldboy to discovery rows where low historical popularity is the main constraint.",
        "Primary measure: positive-rating rate among eligible warm-start users; guardrails: minimum support, exposure count and title-level complaint or hide rate if available.",
        "Owner: content strategy with analytics review. Decision window: one controlled campaign cycle.",
    ], "Use two recommendation cards only in the slide deck, each with evidence, owner and test measure.", "State why these titles qualify: high predicted appeal, aligned actual response, sufficient support and comparatively low popularity.", page_break=False)
    add_slide(doc, 8, "Use experiments for uncertain titles and reduce overexposure", [
        "Recommendation 3: A/B test The Incredibles because predicted appeal is 74.45% but actual positive response is 60.75%.",
        "Recommendation 4: reduce default placement of Home Alone, Star Wars Episode I, Twister and The Mask for this segment; test alternative discovery titles instead.",
        "Randomize eligible users, keep placement and timing consistent, and compare positive response with confidence intervals.",
        "Do not generalize beyond warm-start users or claim revenue impact without exposure and commercial outcome data.",
    ], "Use a simple test design: eligible users -> random assignment -> content variants -> measured response.", "Frame reduction as reallocation within one segment, not a global removal decision.", evidence="Home Alone actual positive rate 17.65%; Star Wars Episode I 19.61%; Twister 10.17%; The Mask 5.66% for qualifying Selective Critics pairs.", page_break=True)

    add_slide(doc, 9, "The final submission is reproducible from repository assets", [
        "Repository contains setup guidance, SQL scripts, KNIME documentation, Power BI project files, evidence outputs and the report.",
        "The IEEE-style paper covers Abstract, Keywords, Introduction, Related Work, Methodology, Experiments and Results, Discussion and Recommendations, Conclusion and References.",
        "The paper should remain within 6-10 pages and cite at least five credible sources.",
        "Before submission, confirm clean credentials, relative paths, model version consistency and a fresh dashboard refresh.",
        "Archive the exact inputs and outputs used in the live demo.",
    ], "Use a repository tree with check marks beside reproducibility-critical folders.", "Name the command or workflow sequence needed to rebuild the outputs. Avoid saying only that the repository is complete.", page_break=False)
    add_slide(doc, 10, "The conclusion keeps decisions within the evidence", [
        "What worked: end-to-end integration, out-of-time evaluation, interpretable models and decision-oriented writeback.",
        "What remains limited: historical data, no exposure or revenue outcomes, overlapping clusters and low warm-start TEST coverage.",
        "Next step: run controlled experiments on qualifying opportunities and monitor coverage before expanding use.",
        "PLAN VERSUS ACTUAL: planned ______ hours; actual ______ hours; gap ______ hours.",
        "Team reflection and corrective action: ________________________________________________",
        "Q and A backup: data quality, split dates, leakage controls, K choice, confusion matrix, coverage and recommendation thresholds.",
    ], "End on one summary diagram: evidence -> bounded decision -> experiment -> learning.", "Invite questions on the evidence chain. If asked about business impact, distinguish model performance from causal campaign lift.", page_break=False)
    path = OUT / "Week_4_Final_Dashboard_Decisions_Slide_Content.docx"
    return path, doc


def main():
    outputs = []
    for path, doc in (week1(), week2(), week3(), week4()):
        doc.core_properties.title = path.stem.replace("_", " ")
        doc.core_properties.subject = "MovieLens 10M weekly presentation content"
        doc.core_properties.author = "MovieLens Content Strategy Project Team"
        doc.save(path)
        outputs.append(path)
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
