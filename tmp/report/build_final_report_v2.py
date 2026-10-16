from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import load_workbook
from PIL import Image, ImageDraw


ROOT = Path(r"C:\Users\doubl\Documents\movielens-content-strategy")
TMP = ROOT / "tmp" / "report"
OUT = ROOT / "output"
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "MovieLens_Content_Strategy_Final_Report.docx"

sys.path.insert(0, str(TMP))
import build_final_report as figures_source  # noqa: E402


NAVY = "1F4E78"
PALE_BLUE = "EEF5FA"
LIGHT_GRAY = "D9D9D9"
DARK = "1F1F1F"
MUTED = "5F6B76"


def make_genre_summary(path: Path):
    wb = load_workbook(ROOT / "python_sql_scripts" / "results" / "07_dashboard_views_results.xlsx", read_only=True, data_only=True)
    ws = wb["genre_kpis"]
    records = []
    for genre, count, avg_rating, positive_rate in ws.iter_rows(min_row=2, values_only=True):
        if genre and genre != "(no genres listed)":
            records.append((str(genre), int(count), float(positive_rate)))
    records.sort(key=lambda row: row[1], reverse=True)
    selected = records[:8]
    film_noir = next(row for row in records if row[0] == "Film-Noir")
    if film_noir not in selected:
        selected.append(film_noir)

    im = Image.new("RGB", (1700, 920), "white")
    draw = ImageDraw.Draw(im)
    title_font = figures_source.font(36, True)
    label_font = figures_source.font(22, True)
    value_font = figures_source.font(20)
    small_font = figures_source.font(17)
    draw.text((70, 35), "Genre volume and positive rating rate", font=title_font, fill="#111111")
    draw.text((70, 95), "Top genres by rating volume plus Film-Noir for rate comparison", font=value_font, fill="#555555")
    max_count = max(row[1] for row in selected)
    y = 165
    for genre, count, rate in selected:
        draw.text((70, y + 9), genre, font=label_font, fill="#222222")
        x0, x1 = 350, 1220
        bar_end = x0 + int((count / max_count) * (x1 - x0))
        draw.rounded_rectangle((x0, y, x1, y + 46), radius=9, fill="#E8EEF4")
        draw.rounded_rectangle((x0, y, bar_end, y + 46), radius=9, fill="#2F75B5")
        draw.text((1245, y + 6), f"{count/1_000_000:.2f}M ratings", font=value_font, fill="#333333")
        draw.text((1490, y + 6), f"{rate:.1%} positive", font=value_font, fill="#548235" if rate >= 0.55 else "#555555")
        y += 74
    draw.text((70, 850), "Volume and positive rate must be interpreted together; a high rate with low support is not sufficient for promotion.", font=small_font, fill="#555555")
    im.save(path, dpi=(220, 220))


def set_run(run, size=10.2, bold=False, italic=False, color=DARK, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def setup_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    ppr = title._element.get_or_add_pPr()
    pborder = ppr.find(qn("w:pBdr"))
    if pborder is not None:
        ppr.remove(pborder)

    for name, size, before, after in (
        ("Heading 1", 15.5, 5, 7),
        ("Heading 2", 12.0, 8, 4),
        ("Heading 3", 10.5, 6, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    cap = doc.styles["Caption"]
    cap.font.name = "Arial"
    cap._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    cap._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    cap.font.size = Pt(8.5)
    cap.font.color.rgb = RGBColor.from_string(DARK)
    cap.font.italic = False
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run("BIG DATA ANALYTICS  |  FINAL PROJECT REPORT"), size=8.2, bold=True, color=MUTED)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_table(doc, headers, rows, widths, font_size=8.8, numeric_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.width = widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(str(value)), size=font_size, bold=True, color="FFFFFF")
    for ridx, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            if ridx % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.width = widths[i]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_run(p.add_run(str(value)), size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = False
    if bold_lead:
        set_run(p.add_run(bold_lead), size=10.2, bold=True)
    set_run(p.add_run(text), size=10.2)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.03
    if not p.runs:
        r = p.add_run(text)
    else:
        p.runs[0].text = text
        r = p.runs[0]
    set_run(r, size=9.8)
    return p


def add_figure(doc, path: Path, caption: str, width: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    doc.add_paragraph(caption, style="Caption")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_run(run, size=8, color=MUTED)


def add_reference(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(f"[{number}] {text}"), size=8.8)


def new_page(doc, heading):
    doc.add_page_break()
    doc.add_heading(heading, level=1)


def build():
    figs = figures_source.make_figures()
    genre_figure = TMP / "figures" / "genre_summary.png"
    make_genre_summary(genre_figure)
    doc = Document()
    setup_document(doc)

    # Cover and required integrity declarations
    p = doc.add_paragraph(style="Title")
    p.add_run("Segment Based Movie Promotion and Content Discovery Strategy")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_run(p.add_run("MovieLens 10M end to end decision support project"), size=14, bold=True, color=NAVY)
    add_table(doc, ["Field", "Details"], [
        ("Course", "Big Data Analytics Final Project 2026"),
        ("Dataset", "MovieLens 10M from GroupLens Research"),
        ("Technology", "MySQL 8.0, KNIME Analytics Platform and Power BI"),
        ("Models", "K Means audience segmentation and Logistic Regression"),
        ("Student name and ID", "________________________________________________________"),
        ("Submission date", "16 October 2026"),
    ], [Cm(4.0), Cm(12.0)], font_size=9.4)
    doc.add_heading("Academic Integrity Declaration", level=2)
    add_body(doc, "By signing below, I certify that the attached assignment is the project team's own work. All datasets, publications, software and AI-assisted work are acknowledged. The team remains responsible for the database execution, KNIME workflows, dashboard, analysis, interpretations and recommendations.")
    p = doc.add_paragraph("Signature __________________________________    Date __________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    for run in p.runs:
        set_run(run, size=9.5)
    doc.add_heading("AI Use Statement", level=2)
    add_body(doc, "Generative AI assisted with planning, technical explanation, code and workflow review, report drafting, formatting and chart preparation. All numerical claims were checked against repository outputs and official documentation. AI-generated text was not treated as independent evidence.")
    doc.add_heading("Abstract", level=2)
    add_body(doc, "This project develops an end-to-end decision-support solution for segment-specific movie promotion using MovieLens 10M. MySQL stores 10,000,054 ratings and derived data products. KNIME reads directly from MySQL, performs reproducible preparation, creates four K Means audience segments and trains a Logistic Regression model for positive ratings, then writes segments, predictions and metrics back to MySQL. The classifier achieved 69.26% accuracy and 75.83% ROC-AUC on 76,633 eligible TEST rows, compared with a 59.32% majority baseline. A five-page Power BI dashboard links these results to controlled promotion, de-prioritisation and testing decisions. The conclusions apply to warm-start rating behaviour, not viewing, revenue or causal campaign impact.")
    p = doc.add_paragraph()
    set_run(p.add_run("Keywords  "), size=9.2, bold=True)
    set_run(p.add_run("MovieLens, audience segmentation, K Means, Logistic Regression, MySQL, KNIME, Power BI, content strategy"), size=9.2, italic=True)

    # 1 Introduction
    new_page(doc, "1 Introduction and Problem Definition")
    add_body(doc, "A movie platform must decide which titles and genres to surface to different users. A single popularity ranking treats all users alike and can repeatedly promote familiar titles even when they fit a segment poorly. The decision-maker in this project is a Content Programming and Personalisation Manager responsible for selecting titles for discovery rows and targeted campaigns.")
    add_body(doc, "The objective is to combine behavioural audience segments with positive-rating probabilities so that the decision-maker can identify titles to promote, maintain, reduce or test. A positive rating is defined as a score of at least 4.0. It is a preference proxy; it is not evidence of a view, purchase, subscription renewal or revenue outcome.")
    doc.add_heading("Decision Questions", level=2)
    add_table(doc, ["Decision question", "Measure", "Dashboard evidence"], [
        ("How do rating behaviour and genre preferences differ across users?", "Segment size, activity, positive rate and genre affinity", "Audience Segments"),
        ("Which user segments are sufficiently interpretable and large enough to support curation?", "K Means profile, segment size and silhouette", "Audience Segments and Model Governance"),
        ("How accurately can prior behaviour identify a positive rating?", "Accuracy, precision, recall, F1, ROC-AUC and confusion matrix", "Positive Engagement Model"),
        ("Which lower-popularity titles merit a controlled promotion test?", "Predicted and actual positive rate, TRAIN popularity and TEST support", "Content Opportunity Matrix"),
    ], [Cm(7.0), Cm(5.3), Cm(4.0)], font_size=8.7)
    add_body(doc, "The required pipeline is implemented as Dataset to MySQL to KNIME to model to MySQL result write-back to interactive dashboard to decision recommendations. Each decision question maps to a stored data product and a dashboard page, which prevents the report from presenting model scores without a decision use.")
    add_body(doc, "The project contributes a reproducible analytical workflow, two interpretable models, five dashboard pages and four evidence-based recommendations. Its main limitation is applicability: the Logistic Regression feature view covers only 5.11% of chronological TEST rows because both the user and movie require TRAIN history.")

    # 2 Dataset
    new_page(doc, "2 Dataset Description")
    add_body(doc, "MovieLens 10M is a public research dataset released by GroupLens Research in January 2009. The README permits research use subject to its stated conditions and identifies GroupLens as the source and owner of the collection [1]. User identifiers are anonymised, and the dataset contains no names, contact details or demographic attributes.")
    add_table(doc, ["Entity", "Rows", "Key fields", "Meaning and analytical role"], [
        ("ratings", "10,000,054", "user_id, movie_id, rating, timestamp", "Explicit preference event on a 0.5 to 5.0 scale"),
        ("movies", "10,681", "movie_id, title, genres", "Movie catalogue and pipe-separated genre labels"),
        ("tags", "95,580", "user_id, movie_id, tag, timestamp", "Free-text annotations supplied by users"),
        ("movie_genres", "21,564", "movie_id, genre", "Normalised bridge used for filtering and aggregation"),
        ("rating users", "69,878", "user_id", "Users represented in the ratings table"),
    ], [Cm(2.1), Cm(2.0), Cm(5.2), Cm(7.0)], font_size=8.5, numeric_cols=(1,))
    add_body(doc, "The GroupLens README reports 71,567 total users because it also includes users represented only in the tags file. The project database contains 69,878 users with at least one rating. Every rating user has at least 20 ratings. The rating timestamps cover 9 January 1995 to 5 January 2009.")
    add_body(doc, "The dataset is relevant because it combines a large interaction table, reusable movie dimensions, multi-value genres and timestamps suitable for future-facing validation. It exceeds the recommended project size by several orders of magnitude and exercises the required database, KNIME and dashboard architecture. However, it does not contain impressions, watch time, licensing cost, revenue or contemporaneous market conditions. Recommendations therefore concern rating preference and controlled discovery tests.")
    add_figure(doc, figs["ratings"], "Figure 1  Distribution of all 10,000,054 MovieLens ratings", 5.7)

    # 3 Data quality
    new_page(doc, "3 Data Quality Assessment")
    add_body(doc, "The project checks completeness, uniqueness, validity, consistency and plausible coverage in both SQL and the reproducible KNIME data-quality workflow. The checks verify referential integrity, allowed rating values, timestamp validity, duplicate user-movie pairs, minimum activity and catalogue coverage. The workflow preserves the original rows and records derived transformations separately.")
    add_table(doc, ["Quality check", "Before cleaning", "After cleaning", "Treatment and rationale"], [
        ("Rating rows", "10,000,054", "10,000,054", "No deletion or imputation was required"),
        ("Missing or orphan ratings", "0", "0", "Foreign-key consistency confirmed"),
        ("Missing or orphan tags", "0", "0", "Foreign-key consistency confirmed"),
        ("Invalid ratings or timestamps", "0", "0", "Allowed range and positive timestamp rules passed"),
        ("Duplicate user-movie ratings", "No rows returned", "No rows returned", "No duplicate removal was required"),
        ("Users below 20 ratings", "0", "0", "Minimum activity condition already satisfied"),
        ("Catalogue movies with ratings", "10,677 of 10,681", "10,677 of 10,681", "Four unrated titles retained as valid catalogue records"),
    ], [Cm(4.1), Cm(3.0), Cm(3.0), Cm(6.2)], font_size=8.2)
    add_body(doc, "Because the checks found no invalid interaction rows, before and after counts are identical. Cleaning consisted of reproducible structural transformations rather than destructive row removal: parsing genres into movie_genres, assigning chronological splits, creating TRAIN-only user and movie histories, generating user-genre profiles and materialising dashboard marts. Low-support titles remain in storage and are controlled through minimum-support filters at decision time.")
    add_body(doc, "The remaining quality risks are conceptual. Ratings are observed only after a user chooses to rate, so missing ratings do not mean dislike. The population overrepresents active raters, and the collection ends in 2009. Genre labels can be broad or multi-valued, while tags are sparse relative to ratings. These limitations affect generalisation even though the stored rows pass integrity checks.")
    add_figure(doc, ROOT / "screenshots" / "knime" / "Knime_WorkFlow_movielens_data_quality.png", "Figure 2  KNIME data quality workflow reading directly from MySQL", 6.1)

    # 4 Architecture page one
    new_page(doc, "4 System Architecture and Data Pipeline")
    add_body(doc, "MySQL is the governed storage layer for source records, derived features and model results. KNIME reads directly from MySQL, performs data preparation and modelling, and writes cluster labels, predictions and metrics back to database tables. Power BI reads compact MySQL views instead of importing the full ratings table. This design preserves a traceable handoff between each layer and keeps the dashboard responsive.")
    add_figure(doc, figs["architecture"], "Figure 3  Implemented architecture from source data to decisions", 6.7)
    doc.add_heading("4.1 MySQL Database and SQL", level=2)
    add_body(doc, "The schema separates source entities from derived and reporting tables. Primary keys protect movie and interaction identity; indexes on user, movie and timestamp support joins, filtering and chronological aggregation. Runnable scripts create the schema, load the delimited files, execute quality checks, create the 70/15/15 split, build TRAIN statistics, assemble Logistic Regression features and publish dashboard views. Filtering, GROUP BY aggregation and joins produce the report KPIs and data-quality evidence.")
    add_table(doc, ["Layer", "Examples", "Purpose"], [
        ("Source", "ratings, movies, tags, movie_genres", "Preserve the documented MovieLens entities"),
        ("Derived", "rating_split, train_user_stats, train_movie_stats, ml_rating_features", "Create reproducible chronological and behavioural features"),
        ("Model results", "user_segments, rating_predictions, model_metrics", "Store KNIME-written labels, probabilities and evaluation"),
        ("Reporting", "dashboard views", "Expose compact and version-aware evidence to Power BI"),
    ], [Cm(2.6), Cm(7.4), Cm(6.3)], font_size=8.5)

    # 4 Architecture page two
    doc.add_page_break()
    doc.add_heading("4.2 KNIME Processing and Result Write Back", level=2)
    add_body(doc, "Three KNIME workflows cover data quality, audience segmentation and positive-engagement prediction. Database Connector and Reader nodes retrieve governed tables. Labelled workflow sections apply type handling, transformations, validation and modelling. Database Writer operations persist the final user segments, rating-level predictions and model metrics to MySQL. This write-back is mandatory for the project and avoids a CSV-only analytical path.")
    add_figure(doc, ROOT / "screenshots" / "knime" / "Knime_WorkFlow_movielens_like_prediction.png", "Figure 4  KNIME Logistic Regression workflow with chronological branches and MySQL write back", 6.6)
    doc.add_heading("Chronological Split and Leakage Controls", level=2)
    add_table(doc, ["Split", "Rows", "Period", "Use"], [
        ("TRAIN", "7,000,042", "1995-01-09 to 2005-04-02", "Feature histories and model fitting"),
        ("VALID", "1,500,006", "2005-04-02 to 2006-12-28", "Model and threshold checks"),
        ("TEST", "1,500,006", "2006-12-28 to 2009-01-05", "Final out-of-time evaluation"),
    ], [Cm(2.0), Cm(2.7), Cm(5.2), Cm(6.4)], font_size=8.7, numeric_cols=(1,))
    add_body(doc, "All user, movie and genre histories are calculated from TRAIN. TRAIN model rows use leave-one-out aggregates so a rating cannot enter its own predictors. VALID and TEST use frozen TRAIN aggregates. user_id and movie_id remain identifiers rather than numerical predictors. The eligible feature view requires both the user and movie to have TRAIN history, which protects feature meaning but reduces TEST coverage.")

    # 5 EDA
    new_page(doc, "5 Exploratory Data Analysis")
    add_body(doc, "The average rating is 3.512, and 50.06% of ratings meet the positive threshold of 4.0 or higher. Rating 4.0 is the most common value with 2,875,850 records. These values indicate that the binary target is not severely imbalanced at the full-dataset level, although temporal and warm-start filtering changes the composition of the model sample.")
    add_table(doc, ["KPI", "Value", "Interpretation"], [
        ("Ratings", "10,000,054", "Large interaction base for aggregation and temporal validation"),
        ("Rating users", "69,878", "Users with explicit rating history"),
        ("Movies", "10,681", "Catalogue size available for content analysis"),
        ("Average rating", "3.512", "Overall preference level"),
        ("Positive rating rate", "50.06%", "Share of ratings at least 4.0"),
        ("Observation window", "1995-01-09 to 2009-01-05", "Historical period and temporal boundary"),
    ], [Cm(4.2), Cm(4.0), Cm(8.1)], font_size=8.8)
    add_body(doc, "Drama has the largest genre rating volume at 4,344,198 and a 56.52% positive rate. Film-Noir has a 71.95% positive rate but much lower support. This contrast shows why a high positive rate alone is not a sufficient promotion rule. Decisions must consider segment fit, historical popularity and the number of eligible TEST observations.")
    add_body(doc, "User activity is highly uneven. Selective Critics later emerge as the most active segment, averaging 385.3 TRAIN ratings per user, while the other final segments average approximately 61 to 67. This difference affects how much title-level support each segment contributes to the content opportunity matrix.")
    add_figure(doc, genre_figure, "Figure 5  Genre rating volume and positive-rating rate", 6.15)

    # 6 ML KMeans
    new_page(doc, "6 Machine Learning")
    doc.add_heading("6.1 K Means Audience Segmentation", level=2)
    add_body(doc, "K Means answers the decision question about interpretable audience groups. The KNIME workflow joins TRAIN user statistics with pivoted genre averages, converts genre ratings to user-relative affinity features, handles missing values and numerical outliers, standardises the feature set and evaluates K values from 3 through 7. The final cluster labels are written to user_segments with a model version and run date.")
    add_table(doc, ["K", "Mean silhouette", "Decision"], [
        ("3", "0.060209", "Fewer profiles but weaker separation"),
        ("4", "0.068639", "Selected for interpretability"),
        ("5", "0.069188", "Only 0.00055 above K equals 4"),
        ("6", "0.064601", "Lower score"),
        ("7", "0.065112", "Lower score"),
    ], [Cm(2.0), Cm(3.8), Cm(10.5)], font_size=8.8, numeric_cols=(1,))
    add_body(doc, "K equals 4 was retained because K equals 5 improved mean silhouette by only 0.00055 while adding another business profile. The absolute silhouette scores remain low, so the segments overlap. They should be treated as descriptive decision lenses rather than natural or demographic categories.")
    add_figure(doc, figs["clusters"], "Figure 6  K selection and the four final audience profiles", 6.35)
    add_table(doc, ["Segment", "Users", "Positive rate", "Profile interpretation"], [
        ("Diverse Explorers", "12,182", "51.38%", "Relative thriller, action and mystery affinity"),
        ("Selective Critics", "10,888", "50.88%", "Highest activity and longest observed history"),
        ("Mainstream Enthusiasts", "14,457", "42.06%", "Relative family and animation affinity"),
        ("Genre Specialists", "17,346", "71.43%", "Largest segment and highest positive tendency"),
    ], [Cm(4.0), Cm(2.2), Cm(2.7), Cm(7.4)], font_size=8.4, numeric_cols=(1, 2))

    # 6 ML Logistic
    doc.add_page_break()
    doc.add_heading("6.2 Logistic Regression for Positive Engagement", level=2)
    add_body(doc, "Logistic Regression estimates the probability that a rating is at least 4.0. Predictors describe prior user behaviour, prior movie behaviour, activity, cluster membership and genre context derived from TRAIN. KNIME fits missing-value and standardisation preprocessing on TRAIN and applies the frozen transformations to VALID and TEST. A 0.5 threshold converts probability into class, while SQL recomputes final metrics from the stored TEST predictions.")
    add_table(doc, ["Metric", "Result", "Interpretation"], [
        ("Accuracy", "69.26%", "9.93 percentage points above the majority baseline"),
        ("Majority baseline", "59.32%", "Reference accuracy for the dominant TEST class"),
        ("Precision", "60.58%", "Share of predicted positives that are positive"),
        ("Recall", "69.94%", "Share of actual positives identified"),
        ("F1", "64.92%", "Balance between precision and recall"),
        ("ROC-AUC", "75.83%", "Ability to rank positive outcomes above negative outcomes"),
    ], [Cm(3.4), Cm(3.0), Cm(9.9)], font_size=8.8, numeric_cols=(1,))
    add_figure(doc, figs["model"], "Figure 7  Logistic Regression TEST metrics and confusion matrix", 6.4)
    add_body(doc, "The confusion matrix contains 31,273 true negatives, 14,189 false positives, 9,369 false negatives and 21,802 true positives. The model improves on the baseline and provides useful ranking information for covered rows, but it is not universally applicable. Only 76,633 of 1,500,006 TEST rows have complete warm-start features, giving 5.11% coverage. The reported metrics therefore describe this eligible subset, not the entire future interaction population.")
    add_body(doc, "Potential leakage is controlled through chronological splitting, TRAIN-only aggregates, leave-one-out TRAIN features and frozen transformations. Remaining risks include temporal drift, unobserved exposure, selection bias among active raters, overlapping clusters and untested probability calibration.")

    # 7 Dashboard
    new_page(doc, "7 Dashboard and Insights")
    add_body(doc, "The Power BI project imports compact MySQL views that include source KPIs and KNIME-written segments, predictions and metrics. Five pages separate descriptive analysis, model interpretation, opportunity selection and governance. Global controls for model version, segment, genre and minimum support prevent incompatible runs or weak cells from being combined.")
    add_figure(doc, figs["dashboard"], "Figure 8  Implemented five-page Power BI report structure from the PBIP definitions", 6.55)
    add_table(doc, ["Dashboard page", "Question answered", "Key visuals and controls"], [
        ("Executive Overview", "What is the scale and overall rating behaviour?", "KPI cards, rating distribution, genre KPIs and top movies"),
        ("Audience Segments", "How do the four groups differ?", "Segment size, activity versus positive rate, profiles and genre heatmap"),
        ("Positive Engagement Model", "How well does the classifier perform?", "Baseline comparison, metrics, coefficients and confusion matrix"),
        ("Content Opportunity Matrix", "Which titles should be promoted, reduced or tested?", "Popularity versus predicted appeal, candidate table and support filter"),
        ("Model Governance", "When are results safe to interpret?", "Split dates, versions, silhouette, coverage and limitations"),
    ], [Cm(4.0), Cm(5.5), Cm(6.8)], font_size=8.4)
    add_body(doc, "The dashboard answers the first two decision questions through segment profiles and K selection, the third through out-of-time classification metrics, and the fourth through the opportunity matrix. Only Selective Critics produced title-segment cells with TEST support of at least 50. The other segments had maximum support between 18 and 30, so their title-level results are exploratory rather than action-ready.")
    add_body(doc, "The opportunity matrix places TRAIN popularity on a logarithmic x-axis, predicted positive probability on the y-axis and TEST support in bubble size. Median reference lines define four action regions. A single model-version filter is required because combining prediction runs would double counts and invalidate support.")

    # 8 Recommendations
    new_page(doc, "8 Decision Recommendations")
    add_body(doc, "The recommendations below answer the decision questions while keeping the action within the available evidence. Each action specifies its empirical basis, expected benefit and principal trade-off.")
    add_table(doc, ["Recommended action", "Supporting evidence", "Expected benefit and trade off"], [
        ("Run a Selective Critics discovery test led by Hotel Rwanda and My Neighbor Totoro", "Hotel Rwanda: support 108, predicted 76.72%, actual 77.78%, TRAIN popularity 215. My Neighbor Totoro: support 58, predicted 74.87%, actual 74.14%, popularity 487.", "Increase discovery of lower-popularity titles. A controlled experiment is required because rating evidence does not prove campaign lift."),
        ("Add City of God and Oldboy to the same controlled candidate pool", "City of God: support 60, predicted 71.66%, actual 70.00%, popularity 968. Oldboy: support 52, predicted 69.94%, actual 67.31%, popularity 72.", "Broaden the discovery slate with aligned predicted and observed preference. The evidence applies only to eligible warm-start users."),
        ("Reduce default Selective Critics placement for Home Alone, Star Wars Episode I, Twister and The Mask", "Each pair has support above 50 but low actual positive rates: 17.65%, 19.61%, 10.17% and 5.66%, respectively, despite high historical popularity.", "Reduce segment-specific mismatch while keeping the titles available elsewhere. Monitor diversity and engagement loss."),
        ("Deploy segment-aware ranking only when warm-start evidence is available", "Classifier TEST coverage is 5.11%; segment write-back covers 54,873 of 69,878 rating users, or 78.5%.", "Avoid unsupported predictions. Use popularity or editorial rules as fallback and track coverage as a deployment KPI."),
    ], [Cm(4.9), Cm(6.2), Cm(5.2)], font_size=8.2)
    doc.add_heading("Evaluation Plan", level=2)
    add_body(doc, "Eligible users should be randomly assigned to the existing experience or the proposed segment-specific slate. Placement, time window and exposure count should remain consistent. Primary outcomes should include click-through, play-start, completion and post-view positive rating when those fields become available. Confidence intervals and minimum sample sizes should be reported before expanding the treatment.")
    add_body(doc, "The Incredibles should be treated separately as a test case rather than an automatic promotion. Its predicted positive rate is 74.45% but its observed rate is 60.75% across 107 qualifying ratings. The disagreement provides a useful diagnostic for model calibration, title context or temporal change.")

    # 9 Conclusion
    new_page(doc, "9 Conclusion and Limitations")
    add_body(doc, "The project implements the required end-to-end decision-support architecture. A documented public dataset is loaded into MySQL; KNIME reads the database, validates and transforms the data, fits K Means and Logistic Regression models, and writes results back; Power BI consumes MySQL views; and recommendations are tied to model evidence, observed response, popularity and support.")
    add_body(doc, "Four audience profiles provide interpretable curation lenses, although their low silhouette values show substantial overlap. The Logistic Regression model improves accuracy by 9.93 percentage points over the majority baseline and achieves 75.83% ROC-AUC on eligible TEST rows. The strongest immediate application is a controlled discovery test for Selective Critics, where title-level support is sufficient to compare predicted and observed positive response.")
    doc.add_heading("Limitations", level=2)
    for item in [
        "MovieLens 10M ends in January 2009 and may not reflect current catalogue, platform or audience behaviour.",
        "Ratings reveal expressed preference after a rating event, not exposure, viewing, revenue, churn or causal campaign impact.",
        "The minimum of 20 ratings per rating user overrepresents active users and limits cold-start relevance.",
        "The Logistic Regression model covers only 5.11% of chronological TEST rows because user and movie histories must exist in TRAIN.",
        "K Means silhouette scores are low in absolute terms, and segment names are interpretive summaries rather than demographic identities.",
        "High-support title-level evidence is concentrated in Selective Critics; other segment recommendations require more observations.",
    ]:
        add_bullet(doc, item)
    doc.add_heading("Potential Improvements", level=2)
    add_body(doc, "Future work should add cold-start features, probability calibration, ranking metrics, stronger rolling temporal validation and exposure-aware outcomes. Modern viewing, impression, completion, licensing-cost and revenue data would support decisions beyond stated preference. Online experiments should estimate causal effects before any broad promotion policy is adopted.")
    add_body(doc, "The final decision is therefore bounded: use the current outputs to prioritise transparent, measurable experiments for warm-start users, while keeping editorial or popularity fallbacks for uncovered cases.")

    # References
    new_page(doc, "References")
    add_reference(doc, 1, "GroupLens Research, MovieLens 10M Dataset and README, 2009. https://grouplens.org/datasets/movielens/10m/ and https://files.grouplens.org/datasets/movielens/ml-10m-README.html. Accessed 16 October 2026.")
    add_reference(doc, 2, "F. M. Harper and J. A. Konstan, The MovieLens Datasets History and Context, ACM Transactions on Interactive Intelligent Systems, vol. 5, no. 4, article 19, 2015. doi:10.1145/2827872.")
    add_reference(doc, 3, "S. P. Lloyd, Least Squares Quantization in PCM, IEEE Transactions on Information Theory, vol. 28, no. 2, pp. 129-137, 1982. doi:10.1109/TIT.1982.1056489.")
    add_reference(doc, 4, "D. W. Hosmer, S. Lemeshow and R. X. Sturdivant, Applied Logistic Regression, 3rd ed., Wiley, 2013.")
    add_reference(doc, 5, "T. Fawcett, An Introduction to ROC Analysis, Pattern Recognition Letters, vol. 27, no. 8, pp. 861-874, 2006. doi:10.1016/j.patrec.2005.10.010.")
    add_reference(doc, 6, "G. James, D. Witten, T. Hastie and R. Tibshirani, An Introduction to Statistical Learning, 2nd ed., Springer, 2021. doi:10.1007/978-1-0716-1418-1.")
    add_reference(doc, 7, "KNIME AG, Connect to a Database Using a Dedicated Connector Node, KNIME Analytics Platform Documentation. https://docs.knime.com/ap/latest/databases/connect-dedicated. Accessed 16 October 2026.")
    add_reference(doc, 8, "KNIME AG, Insert Update or Merge Data, KNIME Analytics Platform Documentation. https://docs.knime.com/ap/latest/databases/insert-update-merge. Accessed 16 October 2026.")
    add_reference(doc, 9, "Oracle, MySQL 8.0 Reference Manual, Loading Data into a Table. https://dev.mysql.com/doc/refman/8.0/en/loading-tables.html. Accessed 16 October 2026.")
    add_reference(doc, 10, "Microsoft, Data Sources in Power BI Desktop. https://learn.microsoft.com/power-bi/connect-data/desktop-data-sources. Accessed 16 October 2026.")

    # Appendices
    new_page(doc, "Appendix A Reproducibility Guide")
    add_body(doc, "Credentials must be stored in the ignored python_sql_scripts .env file or in KNIME credential configuration. No password should be committed to the repository.")
    steps = [
        "Download MovieLens 10M from GroupLens and verify the source files.",
        "Run python_sql_scripts 01 through 05 to create the schema, load source data, execute quality checks, create chronological splits and build TRAIN feature tables.",
        "Open knime/movielens_audience_segmentation.knwf, configure the MySQL Connector, evaluate K values, retain K equals 4 and write user_segments.",
        "Run python_sql_scripts/06_logistic_features.py to create ml_rating_features.",
        "Open knime/movielens_like_prediction.knwf, execute TRAIN, VALID and TEST, and write rating_predictions and model_metrics.",
        "Run python_sql_scripts/07_dashboard_views.py and refresh dashboard/movielens_content_strategy.pbix or the PBIP project.",
        "Confirm split dates, row counts, model version and final metrics against this report before submission.",
    ]
    for step in steps:
        add_bullet(doc, step)
    add_table(doc, ["Required deliverable", "Repository artifact"], [
        ("SQL schema, import and queries", "sql/01_create_schema.sql through sql/07_dashboard_views.sql"),
        ("KNIME workflow export", "knime/movielens_data_quality_finalterm.knwf, movielens_audience_segmentation.knwf and movielens_like_prediction.knwf"),
        ("Interactive dashboard", "dashboard/movielens_content_strategy.pbix and PBIP project"),
        ("Setup and dependencies", "README.md plus python_sql_scripts/README.md and requirements.txt"),
        ("Final report", "output/MovieLens_Content_Strategy_Final_Report.docx; export the approved copy to PDF"),
    ], [Cm(5.0), Cm(11.3)], font_size=8.7)

    new_page(doc, "Appendix B Technical Evidence")
    add_figure(doc, ROOT / "screenshots" / "knime" / "Knime_WorkFlow_movielens_audience_segment.png", "Figure B1  KNIME audience segmentation workflow", 6.55)
    add_table(doc, ["Result table", "Producer", "Core fields", "Dashboard use"], [
        ("user_segments", "KNIME K Means", "user_id, cluster, segment label, model version", "Segment profiles and filters"),
        ("rating_predictions", "KNIME Logistic Regression", "probability, predicted class, actual class, split, version", "Metrics and opportunity matrix"),
        ("model_metrics", "KNIME and SQL", "metric, value, split, run metadata", "Model performance and governance"),
    ], [Cm(3.3), Cm(3.5), Cm(5.7), Cm(3.8)], font_size=8.4)
    add_body(doc, "The report figures are derived from repository outputs, including the quality-check workbook, K Means CSV results, stored Logistic Regression predictions, MySQL dashboard views and PBIP page definitions. The source artifacts remain the authoritative evidence for reproduction.")

    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        p = section.footer.paragraphs[0]
        p.clear()
        add_page_number(p)

    doc.core_properties.title = "Segment Based Movie Promotion and Content Discovery Strategy"
    doc.core_properties.subject = "Big Data Analytics Final Project 2026"
    doc.core_properties.author = "MovieLens Content Strategy Project Team"
    doc.core_properties.keywords = "MovieLens, MySQL, KNIME, Power BI, K Means, Logistic Regression"
    doc.core_properties.comments = "Final report aligned with the Big Data Analytics Final Project 2026 requirements"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
