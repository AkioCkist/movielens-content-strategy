from __future__ import annotations

import csv
import math
from collections import Counter, defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\doubl\Documents\movielens-content-strategy")
TMP = ROOT / "tmp" / "report"
FIG = TMP / "figures"
OUT = ROOT / "output"
FIG.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT_DOCX = OUT / "MovieLens_Content_Strategy_Final_Report.docx"


NAVY = "1F4E78"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9D9D9"
DARK = "1F1F1F"
MUTED = "666666"
GREEN = "548235"
RED = "C65911"


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


F18 = font(18)
F20B = font(20, True)
F24B = font(24, True)
F28B = font(28, True)
F32B = font(32, True)
F36B = font(36, True)


def draw_centered(draw: ImageDraw.ImageDraw, box, text, used_font, fill):
    x0, y0, x1, y1 = box
    lines = text.split("\n")
    bbox = draw.multiline_textbbox((0, 0), text, font=used_font, spacing=5, align="center")
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    draw.multiline_text(
        ((x0 + x1 - tw) / 2, (y0 + y1 - th) / 2),
        text,
        font=used_font,
        fill=fill,
        spacing=5,
        align="center",
    )


def make_architecture(path: Path) -> None:
    im = Image.new("RGB", (1800, 720), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 35), "End to end decision support pipeline", font=F36B, fill="#111111")
    labels = [
        ("MovieLens 10M", "ratings movies tags"),
        ("MySQL", "schema quality checks time split"),
        ("KNIME", "features K Means logistic regression"),
        ("MySQL", "segments predictions metrics marts"),
        ("Power BI", "five interactive pages"),
        ("Decisions", "promote maintain reduce test"),
    ]
    colors = ["#EAF2F8", "#D9EAF7", "#E8F2E4", "#D9EAF7", "#F7E7CE", "#EDE2F3"]
    x = 60
    y = 235
    w = 245
    h = 230
    for i, ((title, sub), color) in enumerate(zip(labels, colors)):
        d.rounded_rectangle((x, y, x + w, y + h), radius=24, fill=color, outline="#5B6573", width=3)
        draw_centered(d, (x + 12, y + 30, x + w - 12, y + 105), title, F28B, "#111111")
        draw_centered(d, (x + 16, y + 105, x + w - 16, y + h - 20), sub, F18, "#333333")
        if i < len(labels) - 1:
            ax0 = x + w + 10
            ax1 = x + w + 55
            ay = y + h // 2
            d.line((ax0, ay, ax1, ay), fill="#355C7D", width=8)
            d.polygon([(ax1, ay), (ax1 - 20, ay - 14), (ax1 - 20, ay + 14)], fill="#355C7D")
        x += 295
    d.text((68, 550), "Reproducibility controls", font=F24B, fill="#111111")
    controls = [
        "versioned SQL scripts",
        "chronological split",
        "TRAIN only aggregates",
        "model version and run time",
        "KNIME write back",
    ]
    x = 65
    for text in controls:
        tw = d.textbbox((0, 0), text, font=F18)[2]
        d.rounded_rectangle((x, 600, x + tw + 36, 655), radius=13, fill="#F2F2F2", outline="#B7B7B7", width=2)
        d.text((x + 18, 615), text, font=F18, fill="#333333")
        x += tw + 56
    im.save(path, dpi=(220, 220))


def make_rating_distribution(path: Path) -> None:
    wb = load_workbook(ROOT / "python_sql_scripts" / "results" / "03_quality_checks_results.xlsx", data_only=True, read_only=True)
    ws = wb["Rating distribution"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    ratings = [str(row[0]) for row in rows]
    counts = [int(row[1]) for row in rows]
    im = Image.new("RGB", (1300, 760), "white")
    d = ImageDraw.Draw(im)
    d.text((65, 40), "Rating distribution", font=F36B, fill="#111111")
    left, top, right, bottom = 105, 135, 1240, 650
    d.line((left, bottom, right, bottom), fill="#555555", width=3)
    d.line((left, top, left, bottom), fill="#555555", width=3)
    maxv = max(counts)
    slot = (right - left) / len(counts)
    barw = slot * 0.64
    for i, (lab, value) in enumerate(zip(ratings, counts)):
        x0 = left + i * slot + (slot - barw) / 2
        x1 = x0 + barw
        y0 = bottom - (value / maxv) * (bottom - top - 35)
        d.rounded_rectangle((x0, y0, x1, bottom), radius=7, fill="#2F75B5")
        d.text((x0 + barw / 2 - 12, bottom + 16), lab, font=F18, fill="#333333")
        label = f"{value / 1_000_000:.2f}M"
        tw = d.textbbox((0, 0), label, font=F18)[2]
        d.text((x0 + (barw - tw) / 2, y0 - 28), label, font=F18, fill="#333333")
    d.text((520, 705), "Rating", font=F20B, fill="#222222")
    d.text((110, 680), "Source: MySQL quality check output", font=F18, fill="#666666")
    im.save(path, dpi=(220, 220))


def read_kmeans_summary():
    import statistics

    results = []
    for k in range(3, 8):
        path = ROOT / "kmeans" / f"{k}.csv"
        values = []
        counts = Counter()
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                values.append(float(row["Silhouette Coefficient"]))
                counts[row["Cluster"]] += 1
        results.append((k, statistics.fmean(values), counts))
    return results


SEGMENTS = [
    ("Diverse Explorers", 12182, 0.5138074, 62.6776),
    ("Selective Critics", 10888, 0.5087552, 385.3217),
    ("Mainstream Enthusiasts", 14457, 0.4206116, 61.2540),
    ("Genre Specialists", 17346, 0.7143489, 66.6189),
]


def make_clustering_summary(path: Path) -> None:
    summary = read_kmeans_summary()
    im = Image.new("RGB", (1700, 850), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 35), "K Means selection and final audience segments", font=F36B, fill="#111111")

    # Left panel: silhouette by K
    left, top, right, bottom = 100, 175, 790, 690
    d.text((100, 105), "Mean silhouette by K", font=F24B, fill="#222222")
    d.line((left, bottom, right, bottom), fill="#555555", width=3)
    d.line((left, top, left, bottom), fill="#555555", width=3)
    values = [v for _, v, _ in summary]
    ymin, ymax = 0.055, 0.072
    points = []
    for i, (k, value, _) in enumerate(summary):
        x = left + 75 + i * 130
        y = bottom - (value - ymin) / (ymax - ymin) * (bottom - top - 20)
        points.append((x, y))
        d.ellipse((x - 10, y - 10, x + 10, y + 10), fill="#2F75B5", outline="#1F4E78")
        d.text((x - 8, bottom + 16), str(k), font=F20B, fill="#333333")
        d.text((x - 36, y - 38), f"{value:.3f}", font=F18, fill="#333333")
    d.line(points, fill="#2F75B5", width=5)
    # Emphasize selected K=4
    x, y = points[1]
    d.ellipse((x - 18, y - 18, x + 18, y + 18), outline="#C65911", width=5)
    d.text((left + 135, 730), "K = 4 retained for interpretability; K = 5 was only 0.00055 higher.", font=F18, fill="#555555")

    # Right panel: segment size and positive rate
    l2, t2, r2, b2 = 920, 175, 1620, 690
    d.text((920, 105), "Final K = 4 profiles", font=F24B, fill="#222222")
    max_users = max(x[1] for x in SEGMENTS)
    colors = ["#5B9BD5", "#A5A5A5", "#ED7D31", "#70AD47"]
    for i, (name, users, rate, activity) in enumerate(SEGMENTS):
        y = t2 + i * 118
        d.text((l2, y), name, font=F20B, fill="#222222")
        barx = l2
        bary = y + 38
        barw = 420 * users / max_users
        d.rounded_rectangle((barx, bary, barx + barw, bary + 36), radius=8, fill=colors[i])
        d.text((barx + barw + 14, bary + 6), f"{users:,} users", font=F18, fill="#333333")
        d.text((l2 + 490, y + 8), f"{rate:.1%} positive", font=F18, fill="#333333")
        d.text((l2 + 490, y + 42), f"{activity:.1f} TRAIN ratings per user", font=F18, fill="#666666")
    d.text((920, 730), "Segmented population: 54,873 users (78.5% of rating users).", font=F18, fill="#555555")
    im.save(path, dpi=(220, 220))


def make_model_evaluation(path: Path) -> None:
    metrics = {
        "Accuracy": 0.69258674,
        "Precision": 0.60576255,
        "Recall": 0.69943216,
        "F1": 0.64923618,
        "AUC": 0.75831740,
        "Baseline": 0.59324312,
    }
    cm = [[31273, 14189], [9369, 21802]]  # rows actual 0/1, cols predicted 0/1
    im = Image.new("RGB", (1700, 820), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 35), "Logistic Regression test evaluation", font=F36B, fill="#111111")

    # Metric bars
    d.text((95, 115), "Metrics", font=F24B, fill="#222222")
    left, top = 120, 185
    for i, (name, value) in enumerate(metrics.items()):
        y = top + i * 82
        d.text((left, y + 7), name, font=F20B, fill="#222222")
        x0, x1 = left + 160, left + 650
        d.rounded_rectangle((x0, y, x1, y + 42), radius=10, fill="#E9EFF5")
        color = "#70AD47" if name != "Baseline" else "#A5A5A5"
        d.rounded_rectangle((x0, y, x0 + (x1 - x0) * value, y + 42), radius=10, fill=color)
        d.text((x1 + 18, y + 7), f"{value:.1%}", font=F20B, fill="#333333")

    # Confusion matrix
    ox, oy, cell = 1020, 205, 230
    d.text((1050, 115), "Confusion matrix", font=F24B, fill="#222222")
    d.text((ox + 95, oy - 45), "Predicted 0", font=F18, fill="#333333")
    d.text((ox + cell + 85, oy - 45), "Predicted 1", font=F18, fill="#333333")
    d.text((ox - 105, oy + 95), "Actual 0", font=F18, fill="#333333")
    d.text((ox - 105, oy + cell + 95), "Actual 1", font=F18, fill="#333333")
    vmax = max(max(row) for row in cm)
    for r in range(2):
        for c in range(2):
            value = cm[r][c]
            intensity = int(245 - 125 * value / vmax)
            fill = f"#{intensity:02x}{min(245, intensity + 18):02x}{min(255, intensity + 40):02x}"
            x0, y0 = ox + c * cell, oy + r * cell
            d.rectangle((x0, y0, x0 + cell, y0 + cell), fill=fill, outline="#FFFFFF", width=6)
            label = f"{value:,}"
            bbox = d.textbbox((0, 0), label, font=F28B)
            d.text((x0 + (cell - (bbox[2] - bbox[0])) / 2, y0 + 92), label, font=F28B, fill="#111111")
    d.text((1010, 710), "TEST rows evaluated: 76,633", font=F20B, fill="#222222")
    d.text((1010, 748), "Warm start coverage: 5.11% of the chronological TEST split", font=F18, fill="#C65911")
    im.save(path, dpi=(220, 220))


def make_opportunity_scatter(path: Path) -> None:
    rows = []
    with (TMP / "content_opportunities.csv").open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["segment_name"] != "Selective Critics":
                continue
            support = int(row["test_support"])
            if support < 20:
                continue
            rows.append(
                {
                    "title": row["title"],
                    "support": support,
                    "pred": float(row["predicted_like_probability"]),
                    "actual": float(row["actual_positive_rate"]),
                    "pop": int(row["train_popularity"]),
                }
            )
    im = Image.new("RGB", (1700, 940), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 35), "Content opportunity view for Selective Critics", font=F36B, fill="#111111")
    left, top, right, bottom = 145, 150, 1590, 790
    d.line((left, bottom, right, bottom), fill="#555555", width=3)
    d.line((left, top, left, bottom), fill="#555555", width=3)
    xs = [math.log10(max(1, row["pop"])) for row in rows]
    ys = [row["pred"] for row in rows]
    xmin, xmax = min(xs), max(xs)
    ymin, ymax = min(ys), max(ys)
    xmid = sorted(xs)[len(xs) // 2]
    ymid = sorted(ys)[len(ys) // 2]
    def xp(v): return left + 30 + (v - xmin) / (xmax - xmin) * (right - left - 60)
    def yp(v): return bottom - 30 - (v - ymin) / (ymax - ymin) * (bottom - top - 60)
    d.line((xp(xmid), top, xp(xmid), bottom), fill="#A6A6A6", width=2)
    d.line((left, yp(ymid), right, yp(ymid)), fill="#A6A6A6", width=2)
    d.text((left + 30, top + 15), "Hidden opportunities", font=F20B, fill="#548235")
    d.text((right - 280, bottom - 45), "Broad promotion risk", font=F20B, fill="#C65911")
    highlight = {
        "Hotel Rwanda (2004)": "#548235",
        "My Neighbor Totoro (Tonari no Totoro) (1988)": "#548235",
        "Home Alone (1990)": "#C65911",
        "Twister (1996)": "#C65911",
    }
    for row, xv in zip(rows, xs):
        x, y = xp(xv), yp(row["pred"])
        radius = 4 + min(13, row["support"] / 10)
        color = highlight.get(row["title"], "#5B9BD5")
        d.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color, outline="#FFFFFF", width=2)
    labels = [
        ("Hotel Rwanda (2004)", 18, -42),
        ("My Neighbor Totoro (Tonari no Totoro) (1988)", 18, 10),
        ("Home Alone (1990)", 18, -5),
        ("Twister (1996)", 18, 28),
    ]
    by_title = {r["title"]: r for r in rows}
    for title, dx, dy in labels:
        row = by_title[title]
        x, y = xp(math.log10(row["pop"])), yp(row["pred"])
        short_title = title.replace(" (2004)", "").replace(" (1990)", "").replace(" (1996)", "").replace(" (1988)", "")
        short_title = short_title.replace(" (Tonari no Totoro)", "")
        d.text((x + dx, y + dy), short_title, font=F18, fill="#222222")
    for tick in [10, 100, 1000, 10000]:
        if xmin <= math.log10(tick) <= xmax:
            x = xp(math.log10(tick))
            d.line((x, bottom, x, bottom + 10), fill="#555555", width=2)
            d.text((x - 18, bottom + 18), f"{tick:,}", font=F18, fill="#333333")
    for tick in [0.2, 0.4, 0.6, 0.8]:
        if ymin <= tick <= ymax:
            y = yp(tick)
            d.line((left - 10, y, left, y), fill="#555555", width=2)
            d.text((left - 70, y - 12), f"{tick:.0%}", font=F18, fill="#333333")
    d.text((650, 855), "TRAIN popularity on logarithmic scale", font=F20B, fill="#222222")
    d.text((25, 420), "Predicted\nlike\nprobability", font=F20B, fill="#222222", spacing=4)
    d.text((150, 895), "Bubble size represents TEST support; display includes pairs with support of at least 20.", font=F18, fill="#666666")
    im.save(path, dpi=(220, 220))


def make_dashboard_map(path: Path) -> None:
    im = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 35), "Power BI analytical pages", font=F36B, fill="#111111")
    pages = [
        ("1 Executive Overview", "Scale and behaviour", "KPI cards\nRating distribution\nGenre and movie tables", "#D9EAF7"),
        ("2 Audience Segments", "Who differs and how", "Segment sizes\nProfiles and genre matrix\nActivity versus liking", "#E8F2E4"),
        ("3 Engagement Model", "How well predictions rank", "Accuracy precision recall\nF1 AUC baseline\nConfusion matrix", "#F7E7CE"),
        ("4 Content Opportunity", "What to test", "Segment movie scatter\nSupport filter\nCandidate table", "#EDE2F3"),
        ("5 Model Governance", "Can results be trusted", "Split periods\nModel versions\nSilhouette and limitations", "#F2F2F2"),
    ]
    x_positions = [70, 410, 750, 1090, 1430]
    for x, (title, question, visuals, fill) in zip(x_positions, pages):
        d.rounded_rectangle((x, 190, x + 285, 710), radius=22, fill=fill, outline="#6B7280", width=3)
        draw_centered(d, (x + 16, 215, x + 269, 305), title, F24B, "#111111")
        d.line((x + 24, 325, x + 261, 325), fill="#A6A6A6", width=2)
        draw_centered(d, (x + 18, 345, x + 267, 430), question, F20B, "#333333")
        draw_centered(d, (x + 24, 465, x + 261, 650), visuals, F18, "#333333")
    d.text((75, 780), "Shared controls", font=F24B, fill="#111111")
    controls = ["model version", "segment", "genre", "minimum support"]
    x = 295
    for item in controls:
        tw = d.textbbox((0, 0), item, font=F20B)[2]
        d.rounded_rectangle((x, 760, x + tw + 44, 825), radius=14, fill="#FFFFFF", outline="#2F75B5", width=3)
        d.text((x + 22, 780), item, font=F20B, fill="#1F4E78")
        x += tw + 78
    im.save(path, dpi=(220, 220))


def make_figures() -> dict[str, Path]:
    paths = {
        "architecture": FIG / "architecture.png",
        "ratings": FIG / "rating_distribution.png",
        "clusters": FIG / "clustering_summary.png",
        "model": FIG / "model_evaluation.png",
        "opportunities": FIG / "content_opportunity.png",
        "dashboard": FIG / "dashboard_map.png",
    }
    make_architecture(paths["architecture"])
    make_rating_distribution(paths["ratings"])
    make_clustering_summary(paths["clusters"])
    make_model_evaluation(paths["model"])
    make_opportunity_scatter(paths["opportunities"])
    make_dashboard_map(paths["dashboard"])
    return paths


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_columns(section, count: int, space_twips: int = 360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_twips))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run.font.name = "Arial"
    run.font.size = Pt(8)


def set_run_font(run, name="Arial", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(3.5)
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in [
        ("Heading 1", 13, 10, 4),
        ("Heading 2", 10.5, 7, 3),
        ("Heading 3", 9.5, 5, 2),
    ]:
        st = doc.styles[style_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(8)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(DARK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    if "Abstract" not in doc.styles:
        abst = doc.styles.add_style("Abstract", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abst = doc.styles["Abstract"]
    abst.font.name = "Arial"
    abst._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    abst._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    abst.font.size = Pt(9)
    abst.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abst.paragraph_format.left_indent = Cm(0.65)
    abst.paragraph_format.right_indent = Cm(0.65)
    abst.paragraph_format.space_after = Pt(5)


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, size=9.2, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=9.2)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        set_run_font(run, size=9.0)
    if not p.runs:
        set_run_font(p.add_run(text), size=9.0)
    else:
        p.runs[0].text = text
    return p


def add_table(doc: Document, headers, rows, widths=None, font_size=8.0, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_run_font(run, size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = widths[i]
    for rix, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            if rix % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = (alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc: Document, path: Path, caption: str, width: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split("  ", 1)[-1])
    cp = doc.add_paragraph(caption, style="Caption")
    return cp


def heading(doc: Document, text: str, level=1):
    return doc.add_heading(text, level=level)


def add_reference(doc: Document, number: int, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"[{number}] {text}")
    set_run_font(r, size=8.1)


def build_report(figs: dict[str, Path]) -> None:
    doc = Document()
    style_document(doc)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.add_run("Segment Based Movie Promotion and Content Discovery Strategy")
    title_ppr = p._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MovieLens 10M decision support project")
    set_run_font(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(28)

    meta = [
        ("Course", "Big Data Analytics"),
        ("Deliverable", "Final project report and Week 4 research paper"),
        ("Dataset", "MovieLens 10M from GroupLens Research"),
        ("Analytical stack", "MySQL 8.0, KNIME Analytics Platform and Power BI"),
        ("Models", "K Means audience segmentation and Logistic Regression"),
        ("Report date", "16 October 2026"),
    ]
    add_table(
        doc,
        ["Field", "Details"],
        meta,
        widths=[Cm(4.1), Cm(11.7)],
        font_size=9.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Student name and ID")
    set_run_font(r, size=10, bold=True)
    p = doc.add_paragraph("____________________________________________________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading(doc, "Academic Integrity Declaration", 2)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.1)
    p.paragraph_format.right_indent = Cm(1.1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(
        "I certify that this submission represents the project team's work. Sources, software, and AI assistance are acknowledged. "
        "The team remains responsible for the SQL execution, KNIME workflows, model evaluation, dashboard, interpretations, and recommendations."
    )
    set_run_font(r, size=9.5)
    p = doc.add_paragraph("Signature __________________________________    Date __________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)

    # Abstract on cover, as IEEE-style opening content
    heading(doc, "Abstract", 2)
    p = doc.add_paragraph(style="Abstract")
    r = p.add_run(
        "This project builds an end-to-end decision-support pipeline for segment-specific movie promotion using MovieLens 10M. "
        "The solution loads 10,000,054 ratings into MySQL, performs reproducible quality checks and chronological splitting, engineers TRAIN-only behavioural features, and uses KNIME to create four K-Means audience segments and a Logistic Regression model for positive ratings. "
        "The classifier achieved 69.26% accuracy and 75.83% ROC-AUC on 76,633 eligible TEST rows, compared with a 59.32% majority baseline. "
        "The Power BI model converts these outputs into segment profiles and a content-opportunity matrix. The strongest immediate use is controlled curation for warm-start users, not a claim of viewing or revenue uplift."
    )
    set_run_font(r, size=9.0)
    p = doc.add_paragraph(style="Abstract")
    r = p.add_run("Keywords  MovieLens, audience segmentation, K Means, Logistic Regression, MySQL, KNIME, Power BI, content strategy")
    set_run_font(r, size=8.8, italic=True)

    # Main section, two columns
    main = doc.add_section(WD_SECTION.NEW_PAGE)
    main.page_width = Cm(21)
    main.page_height = Cm(29.7)
    main.top_margin = Cm(1.45)
    main.bottom_margin = Cm(1.5)
    main.left_margin = Cm(1.45)
    main.right_margin = Cm(1.45)
    main.footer_distance = Cm(0.65)
    set_columns(main, 2, 420)

    heading(doc, "1 Introduction and Problem Definition", 1)
    add_body(
        doc,
        "A movie platform must decide which titles and genres to surface to different users. A single popularity ranking ignores variation in taste and can over-promote titles that are broadly known but poorly matched to a particular audience. This project supports the Content Programming and Personalization Manager by identifying behavioural segments and estimating whether a warm-start user will give a movie at least four stars. The predicted outcome is positive rating behaviour. It is not a prediction of viewing, revenue, retention, or subscription conversion."
    )
    add_body(
        doc,
        "The central decision is which movies or genres should receive more, equal, or less promotional priority within each segment. The analysis answers four measurable questions."
    )
    add_table(
        doc,
        ["Decision question", "Evidence used"],
        [
            ("How do rating behaviour and genre preferences differ?", "User statistics and segment genre profiles"),
            ("Which audience segments are interpretable and large enough to use?", "K Means profiles, sizes and silhouette"),
            ("What is the probability of a positive rating?", "Logistic Regression probability and TEST metrics"),
            ("Which lower-popularity titles merit a controlled promotion test?", "Segment movie support, TRAIN popularity and predicted probability"),
        ],
        widths=[Cm(3.6), Cm(3.7)],
        font_size=7.6,
    )
    add_body(
        doc,
        "The project follows the required flow Dataset to MySQL to KNIME to MySQL to interactive dashboard to recommendations. The repository contains runnable SQL and Python pipeline scripts, exported KNIME workflows, Power BI project files, query results, and implementation evidence."
    )

    heading(doc, "2 Related Work", 1)
    add_body(
        doc,
        "MovieLens records explicit user preferences as user, item, rating, and timestamp tuples and is widely used for recommendation research [1], [2]. These ratings are useful for modelling stated preferences but do not reveal whether a user was exposed to a title or chose to watch it. The report therefore treats a rating of at least 4.0 as a positive-engagement proxy and keeps business claims within that definition."
    )
    add_body(
        doc,
        "K-Means provides a transparent way to group users around behavioural centroids [3]. Logistic Regression complements segmentation by producing an interpretable probability for a binary outcome. ROC-AUC is used as a ranking measure alongside threshold-dependent metrics such as accuracy, precision, recall, and F1 [8]. KNIME supplies the visual workflow and database nodes used for read and write-back operations [4], [5]."
    )

    heading(doc, "3 Dataset and Data Quality", 1)
    heading(doc, "3.1 Dataset Description", 2)
    add_body(
        doc,
        "MovieLens 10M was released in January 2009 by GroupLens Research. Its README reports 10,000,054 ratings and 95,580 tag applications across 10,681 movies and 71,567 users [1]. The project database contains 69,878 users with at least one rating. The difference is expected because the README user total includes people represented only in the tags file. Every sampled rating user has at least 20 ratings, user identifiers are anonymised, and no demographic fields are supplied."
    )
    add_table(
        doc,
        ["Entity", "Rows", "Role"],
        [
            ("ratings", "10,000,054", "User movie rating and timestamp"),
            ("movies", "10,681", "Title and pipe-separated genres"),
            ("tags", "95,580", "User-generated tag applications"),
            ("movie_genres", "21,564", "Normalised movie genre bridge"),
            ("rating users", "69,878", "Users with at least one rating"),
        ],
        widths=[Cm(2.0), Cm(1.8), Cm(3.5)],
        font_size=7.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_figure(doc, figs["ratings"], "Figure 1  Distribution of the 10,000,054 ratings", 3.15)

    heading(doc, "3.2 Quality Assessment and Cleaning", 2)
    add_body(
        doc,
        "The SQL checks covered completeness, uniqueness, validity, consistency, and plausible coverage. No orphan ratings or tags were found, no invalid rating or zero timestamp was found, no duplicate user-movie pair was returned, and every rating user met the 20-rating minimum. Four catalogue movies had no rating because 10,677 of 10,681 movie IDs appeared in the ratings table. They were retained as valid catalogue records."
    )
    add_table(
        doc,
        ["Check", "Before", "After"],
        [
            ("Rating rows", "10,000,054", "10,000,054"),
            ("Rating range", "0.5 to 5.0", "0.5 to 5.0"),
            ("Orphan ratings", "0", "0"),
            ("Orphan tags", "0", "0"),
            ("Invalid ratings or timestamps", "0", "0"),
            ("Users below 20 ratings", "0", "0"),
        ],
        widths=[Cm(3.6), Cm(1.8), Cm(1.8)],
        font_size=7.7,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(
        doc,
        "The raw data did not require row deletion or imputation for these checks, so the before and after counts are identical. Reproducible transformations created the genre bridge, chronological split, TRAIN-only user and movie aggregates, user-genre profiles, model features, and dashboard marts. Low-support movies were not deleted from storage; minimum-support filters are applied when making decisions."
    )

    # Full-width architecture and KNIME evidence page
    wide = doc.add_section(WD_SECTION.NEW_PAGE)
    wide.page_width = Cm(21)
    wide.page_height = Cm(29.7)
    wide.top_margin = Cm(1.45)
    wide.bottom_margin = Cm(1.5)
    wide.left_margin = Cm(1.6)
    wide.right_margin = Cm(1.6)
    wide.footer_distance = Cm(0.65)
    set_columns(wide, 1)
    heading(doc, "4 Methodology and System Architecture", 1)
    add_body(
        doc,
        "MySQL stores both source data and model outputs. KNIME reads directly from MySQL, performs reproducible transformations and modelling, and writes segments, predictions, and model metrics back to MySQL. Power BI imports compact views rather than the 10-million-row ratings table. This design makes the dashboard responsive and preserves a traceable chain from source rows to recommendations."
    )
    add_figure(doc, figs["architecture"], "Figure 2  Required system architecture and result write-back", 6.55)
    add_figure(
        doc,
        ROOT / "screenshots" / "knime" / "movielens_like_prediction.png",
        "Figure 3  KNIME Logistic Regression workflow with TRAIN VALID TEST branches and MySQL write-back",
        6.7,
    )
    heading(doc, "4.1 Data Products and Ownership", 2)
    add_table(
        doc,
        ["Data product", "Created by", "Primary consumer", "Purpose"],
        [
            ("rating_split", "MySQL step 04", "Feature SQL and governance page", "Chronological TRAIN VALID TEST assignment"),
            ("train user movie and genre statistics", "MySQL step 05", "Both KNIME workflows", "Leakage-controlled behavioural history"),
            ("user_segments", "KNIME K Means", "Logistic features and dashboard", "Segment label version and run date"),
            ("rating_predictions", "KNIME Logistic Regression", "Metric SQL and opportunity view", "TEST probability class and actual label"),
            ("model_metrics", "KNIME and MySQL", "Model and governance pages", "TEST metrics coefficients and silhouette"),
            ("dashboard views", "MySQL step 07", "Power BI", "Compact decision-facing aggregates"),
        ],
        widths=[Cm(4.0), Cm(3.3), Cm(4.7), Cm(4.4)],
        font_size=7.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    # Return to two columns
    main2 = doc.add_section(WD_SECTION.NEW_PAGE)
    main2.page_width = Cm(21)
    main2.page_height = Cm(29.7)
    main2.top_margin = Cm(1.45)
    main2.bottom_margin = Cm(1.5)
    main2.left_margin = Cm(1.45)
    main2.right_margin = Cm(1.45)
    main2.footer_distance = Cm(0.65)
    set_columns(main2, 2, 420)

    heading(doc, "4.2 MySQL Design", 2)
    add_body(
        doc,
        "The schema separates source tables from derived tables. Primary keys protect entity identity, while indexes on rating user, movie, and timestamp support aggregation and chronological access. Bulk loading uses MySQL LOAD DATA, which is appropriate for delimited files [6]. Large derived tables are built in staging tables and atomically renamed, leaving a rollback copy. The Python runner enforces step order and waits for KNIME write-back before continuing."
    )
    add_body(
        doc,
        "The chronological 70/15/15 split sorts by rating timestamp and rating ID. TRAIN contains 7,000,042 rows through 2 April 2005, VALID contains 1,500,006 rows through 28 December 2006, and TEST contains 1,500,006 rows through 5 January 2009. This split is more credible than a random split for a future-facing decision question."
    )
    add_table(
        doc,
        ["Split", "Rows", "Period"],
        [
            ("TRAIN", "7,000,042", "1995-01-09 to 2005-04-02"),
            ("VALID", "1,500,006", "2005-04-02 to 2006-12-28"),
            ("TEST", "1,500,006", "2006-12-28 to 2009-01-05"),
        ],
        widths=[Cm(1.2), Cm(1.7), Cm(4.2)],
        font_size=7.7,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    heading(doc, "4.3 Feature Engineering and Leakage Control", 2)
    add_body(
        doc,
        "User statistics, movie statistics, and user-genre profiles are calculated from TRAIN only. For TRAIN model rows, the row's own rating is removed from user and movie average and positive-rate features. VALID and TEST use frozen TRAIN aggregates. User ID and movie ID remain identifiers and are not treated as continuous predictors. The model includes behavioural histories, activity measures, movie popularity, cluster category, and genre indicators."
    )
    add_body(
        doc,
        "The view requires both the user and movie to have TRAIN history. This warm-start rule preserves feature meaning but sharply reduces future coverage. The feature view retained 6,999,878 TRAIN rows, 244,169 VALID rows, and 76,633 TEST rows. TEST coverage is 5.11%, so the headline metrics apply only to the eligible subset."
    )

    heading(doc, "5 Experiments and Results", 1)
    heading(doc, "5.1 Exploratory Findings", 2)
    add_body(
        doc,
        "The average rating is 3.512 and the positive-rating rate is 50.06%. Ratings of 4.0 are most common at 2,875,850 records. Drama has the largest genre rating volume at 4,344,198 and a 56.52% positive rate. Film-Noir has a 71.95% positive rate but far lower support, illustrating why both preference and popularity are needed when selecting content."
    )

    heading(doc, "5.2 K Means Segmentation", 2)
    add_body(
        doc,
        "The KNIME workflow joins user statistics with pivoted genre averages, converts genre ratings to user-relative affinity values, handles missing values and numerical outliers, standardises model features, and compares K from 3 through 7. Mean silhouette rose from 0.0602 at K=3 to 0.0686 at K=4 and 0.0692 at K=5. K=4 was retained because the 0.00055 gain at K=5 did not justify an additional segment for the business workflow. The low absolute silhouette values show that the groups overlap and should be used as descriptive personas rather than natural categories."
    )
    add_figure(doc, figs["clusters"], "Figure 4  Cluster selection and final K equals 4 segment profiles", 3.15)
    add_table(
        doc,
        ["Segment", "Users", "Positive rate", "Interpretation"],
        [
            ("Genre Specialists", "17,346", "71.43%", "Highest positive tendency and largest segment"),
            ("Mainstream Enthusiasts", "14,457", "42.06%", "Lower positive tendency; family genres stand out relatively"),
            ("Diverse Explorers", "12,182", "51.38%", "Stronger relative affinity for thriller action mystery"),
            ("Selective Critics", "10,888", "50.88%", "Much higher activity and longest observed history"),
        ],
        widths=[Cm(2.1), Cm(1.1), Cm(1.3), Cm(2.8)],
        font_size=7.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    heading(doc, "5.3 Logistic Regression", 2)
    add_body(
        doc,
        "The target is liked equals one when rating is at least 4.0. KNIME trains a Logistic Regression model on a deterministic 20% sample of the eligible TRAIN feature view, fits missing-value and z-score preprocessing on TRAIN, applies the frozen transformations to VALID and TEST, and uses a 0.5 decision threshold. Predictions are written to rating_predictions, while SQL recomputes the final TEST metrics from the stored predictions."
    )
    add_body(
        doc,
        "On 76,633 eligible TEST rows, the model achieved 69.26% accuracy, 60.58% precision, 69.94% recall, 64.92% F1, and 75.83% ROC-AUC. Accuracy exceeds the 59.32% majority baseline by 9.93 percentage points. The confusion matrix contains 31,273 true negatives, 21,802 true positives, 14,189 false positives, and 9,369 false negatives. The model therefore ranks positive outcomes materially better than chance, but its probability calibration and coverage must be considered before use."
    )
    add_figure(doc, figs["model"], "Figure 5  Logistic Regression TEST metrics and confusion matrix", 3.15)

    dashboard = doc.add_section(WD_SECTION.NEW_PAGE)
    dashboard.page_width = Cm(21)
    dashboard.page_height = Cm(29.7)
    dashboard.top_margin = Cm(1.45)
    dashboard.bottom_margin = Cm(1.5)
    dashboard.left_margin = Cm(1.6)
    dashboard.right_margin = Cm(1.6)
    dashboard.footer_distance = Cm(0.65)
    set_columns(dashboard, 1)
    heading(doc, "5.4 Dashboard Implementation", 2)
    add_body(
        doc,
        "The Power BI project is stored as both PBIX and PBIP artifacts. Its semantic model imports compact MySQL views for headline KPIs, rating distribution, genre KPIs, movie popularity, chronological splits, segment profiles, segment-genre statistics, content opportunities, the confusion matrix, and model metrics. Microsoft documents MySQL as a supported Power BI Desktop data source [7]."
    )
    add_figure(doc, figs["dashboard"], "Figure 6  Purpose and evidence on the five Power BI pages", 6.65)
    add_table(
        doc,
        ["Dashboard page", "Decision use", "Primary interaction"],
        [
            ("Executive Overview", "Establish dataset scale and overall rating behaviour", "Top N movie view and genre comparison"),
            ("Audience Segments", "Compare segment size, activity, positive rate, and genre preference", "Segment and genre slicers"),
            ("Positive Engagement Model", "Compare the model with baseline and inspect errors", "Single-select model version"),
            ("Content Opportunity Matrix", "Find lower-popularity titles with high predicted liking", "Model, segment, genre, and support filters"),
            ("Model Governance", "Review split periods, versions, silhouette, and limits", "No incompatible shared model slicer"),
        ],
        widths=[Cm(3.8), Cm(7.0), Cm(5.6)],
        font_size=8.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(
        doc,
        "The opportunity matrix uses TRAIN popularity on a logarithmic x-axis, predicted like probability on the y-axis, and TEST support as bubble size. Median lines separate four decision regions. Because predictions from multiple runs would otherwise double support, every model-result visual must filter a single model_version. The minimum-support control protects decisions from unstable movie-segment cells."
    )

    # Full-width decision evidence
    wide2 = doc.add_section(WD_SECTION.NEW_PAGE)
    wide2.page_width = Cm(21)
    wide2.page_height = Cm(29.7)
    wide2.top_margin = Cm(1.45)
    wide2.bottom_margin = Cm(1.5)
    wide2.left_margin = Cm(1.6)
    wide2.right_margin = Cm(1.6)
    wide2.footer_distance = Cm(0.65)
    set_columns(wide2, 1)
    heading(doc, "6 Dashboard Insights and Decisions", 1)
    add_body(
        doc,
        "The Power BI project contains five pages: Executive Overview, Audience Segments, Positive Engagement Model, Content Opportunity Matrix, and Model Governance. The semantic model reads MySQL views for KPIs, split dates, segment profiles, confusion-matrix counts, model metrics, and segment-movie opportunities. Filters for model version, segment, genre, and minimum support prevent incompatible runs or weak evidence from being mixed."
    )
    add_figure(doc, figs["opportunities"], "Figure 7  Reconstructed content opportunity matrix from the MySQL dashboard view", 6.6)
    add_body(
        doc,
        "Only Selective Critics produced movie-segment pairs with TEST support of at least 50. The other segments had maximum support between 18 and 30 because the chronological warm-start rule left few eligible rows. Operational recommendations should therefore begin with Selective Critics and use lower-support results for exploration only."
    )

    heading(doc, "6.1 Evidence Based Recommendations", 2)
    add_table(
        doc,
        ["Action", "Evidence", "Expected benefit and control"],
        [
            (
                "Test a Selective Critics collection led by Hotel Rwanda and My Neighbor Totoro",
                "Hotel Rwanda: 108 TEST ratings, 76.72% predicted and 77.78% actual positive rate, 215 TRAIN ratings. My Neighbor Totoro: 58 TEST ratings, 74.87% predicted and 74.14% actual positive rate, 487 TRAIN ratings.",
                "Improve discovery of lower-popularity titles. Run an A/B test and track click-through, play-start, completion, and post-view rating before expanding.",
            ),
            (
                "Reduce segment-specific priority for Home Alone, Twister, and The Mask",
                "All have at least 51 TEST ratings and high TRAIN popularity. Predicted positive rates are 27.30%, 37.19%, and 37.74%; observed positive rates are 17.65%, 10.17%, and 5.66%.",
                "Reduce mismatch within Selective Critics while keeping the titles available elsewhere. Monitor engagement loss and catalogue diversity.",
            ),
            (
                "Use segment-aware ranking only for warm-start users",
                "The classifier covers 76,633 of 1,500,006 TEST rows, or 5.11%. Segment write-back covers 54,873 of 69,878 rating users, or 78.5%.",
                "Avoid unsupported predictions. Use popularity or editorial rules as fallback and measure coverage as a deployment KPI.",
            ),
            (
                "Prioritise the Genre Specialists segment for curated pilots",
                "This is the largest segment at 17,346 users and has the highest mean positive rate at 71.43%.",
                "Concentrate initial editorial effort where expressed liking is highest, while testing titles within the segment rather than promoting all genres equally.",
            ),
        ],
        widths=[Cm(4.1), Cm(6.1), Cm(6.2)],
        font_size=8.0,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    # Final two-column section
    main3 = doc.add_section(WD_SECTION.NEW_PAGE)
    main3.page_width = Cm(21)
    main3.page_height = Cm(29.7)
    main3.top_margin = Cm(1.45)
    main3.bottom_margin = Cm(1.5)
    main3.left_margin = Cm(1.45)
    main3.right_margin = Cm(1.45)
    main3.footer_distance = Cm(0.65)
    set_columns(main3, 2, 420)

    heading(doc, "7 Conclusion and Limitations", 1)
    add_body(
        doc,
        "The project meets the required technical pattern: a documented public dataset is loaded into MySQL; KNIME reads the database, prepares features, trains two approved model types, and writes results back; Power BI consumes MySQL views; and recommendations are tied to observed support, model probabilities, and segment profiles. The main decision value is a transparent shortlist for controlled content discovery tests."
    )
    add_body(
        doc,
        "The evidence does not justify claims about viewing, revenue, churn, or causal uplift. MovieLens 10M reflects behaviour collected no later than January 2009 and excludes demographic, exposure, impression, click, watch-time, licensing-cost, and revenue data. Users have at least 20 ratings, so the population overrepresents active raters. Missing ratings cannot distinguish an unseen title from indifference."
    )
    add_body(
        doc,
        "Model coverage is the most important technical limitation. The TRAIN-only warm-start join excludes 94.89% of chronological TEST rows, and high-support movie-level recommendations are concentrated in Selective Critics. Segment separation is also weak in absolute silhouette terms. Future work should add cold-start handling, probability calibration, stronger temporal evaluation, ranking metrics, exposure-aware data, and online experiments."
    )

    heading(doc, "References", 1)
    add_reference(doc, 1, "GroupLens Research, MovieLens 10M Dataset and README, 2009. https://grouplens.org/datasets/movielens/10m/ and https://files.grouplens.org/datasets/movielens/ml-10m-README.html. Accessed 16 October 2026.")
    add_reference(doc, 2, "F. M. Harper and J. A. Konstan, The MovieLens Datasets History and Context, ACM Transactions on Interactive Intelligent Systems, vol. 5, no. 4, article 19, 2015. doi:10.1145/2827872.")
    add_reference(doc, 3, "S. P. Lloyd, Least Squares Quantization in PCM, IEEE Transactions on Information Theory, vol. 28, no. 2, pp. 129-137, 1982. doi:10.1109/TIT.1982.1056489.")
    add_reference(doc, 4, "KNIME AG, Connect to a Database Using a Dedicated Connector Node, KNIME Analytics Platform Documentation. https://docs.knime.com/ap/latest/databases/connect-dedicated. Accessed 16 October 2026.")
    add_reference(doc, 5, "KNIME AG, Insert Update or Merge Data, KNIME Analytics Platform Documentation. https://docs.knime.com/ap/latest/databases/insert-update-merge. Accessed 16 October 2026.")
    add_reference(doc, 6, "Oracle, MySQL 8.0 Reference Manual, Loading Data into a Table. https://dev.mysql.com/doc/refman/8.0/en/loading-tables.html. Accessed 16 October 2026.")
    add_reference(doc, 7, "Microsoft, Data Sources in Power BI Desktop. https://learn.microsoft.com/power-bi/connect-data/desktop-data-sources. Accessed 16 October 2026.")
    add_reference(doc, 8, "T. Fawcett, An Introduction to ROC Analysis, Pattern Recognition Letters, vol. 27, no. 8, pp. 861-874, 2006. doi:10.1016/j.patrec.2005.10.010.")
    add_reference(doc, 9, "G. James, D. Witten, T. Hastie, and R. Tibshirani, An Introduction to Statistical Learning, 2nd ed., Springer, 2021. doi:10.1007/978-1-0716-1418-1.")

    heading(doc, "Appendix A Reproducibility Guide", 1)
    add_body(
        doc,
        "A reviewer can reproduce the pipeline in the following order. Credentials belong in the ignored python_sql_scripts .env file or in KNIME credential configuration and must never be committed."
    )
    steps = [
        "Download MovieLens 10M from GroupLens and verify the supplied checksum.",
        "Run SQL or Python steps 01 through 05 to create the schema, load data, run checks, create the split, and build TRAIN features.",
        "Open movielens_audience_segmentation.knwf, configure the MySQL Connector, execute K-Means evaluation, retain K equals 4, and write user_segments.",
        "Run step 06 to create ml_rating_features.",
        "Open movielens_like_prediction.knwf, execute TRAIN and VALID, freeze the threshold, execute TEST, and write rating_predictions and model_metrics.",
        "Run step 07 to create dashboard views, then refresh movielens_content_strategy.pbix or the PBIP project.",
        "Confirm row counts, model version logit_v1, split dates, metrics, and recommendation evidence against the report.",
    ]
    for step in steps:
        add_bullet(doc, step)

    heading(doc, "Appendix B AI Use Statement", 1)
    add_body(
        doc,
        "Generative AI was used to support project planning, explain technical concepts, review SQL and workflow structure, draft and format this report, and prepare charts from verified project outputs. Dataset selection, SQL execution, KNIME workflow construction, model evaluation, dashboard development, interpretations, and recommendations remain the responsibility of the project team. AI-generated wording and suggestions were checked against repository files, database results, and official documentation and were not treated as independent evidence."
    )

    appendix = doc.add_section(WD_SECTION.NEW_PAGE)
    appendix.page_width = Cm(21)
    appendix.page_height = Cm(29.7)
    appendix.top_margin = Cm(1.55)
    appendix.bottom_margin = Cm(1.55)
    appendix.left_margin = Cm(1.7)
    appendix.right_margin = Cm(1.7)
    appendix.footer_distance = Cm(0.65)
    set_columns(appendix, 1)
    heading(doc, "Appendix C Weekly Presentation Alignment", 1)
    add_body(
        doc,
        "The repository now contains the evidence required for the four weekly presentations. The table below maps each week's requested focus to the artifacts and results available for demonstration."
    )
    add_table(
        doc,
        ["Week", "Required focus", "Project evidence", "Presentation emphasis"],
        [
            ("1", "Problem, dataset, decision questions, architecture, plan", "Project plan, four decision questions, MovieLens source, required pipeline", "Explain the stakeholder decision and why ratings support only preference claims"),
            ("2", "Quality, MySQL, KNIME integration, early KPI", "Schema and load scripts, quality-check workbook, split and feature tables, SQL screenshots", "Show zero invalid or orphan rows and direct database access"),
            ("3", "Features, model evaluation, write-back, dashboard draft", "Two KNIME workflows, K files, TEST metrics, user_segments, rating_predictions, model_metrics", "Lead with leakage control, K selection, baseline comparison, and write-back"),
            ("4", "Final dashboard, decisions, report, plan versus actual", "Five-page PBIX and PBIP, recommendation evidence, reproducibility guide, final DOCX", "State warm-start coverage and present recommendations as controlled tests"),
        ],
        widths=[Cm(1.0), Cm(4.3), Cm(5.9), Cm(5.4)],
        font_size=8.4,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    heading(doc, "Submission Readiness", 2)
    add_table(
        doc,
        ["Deliverable", "Available artifact", "Final check"],
        [
            ("SQL and schema", "sql 01 through 07 plus knime_logit_metrics", "Run from a clean schema or retain exported result workbooks"),
            ("KNIME", "Audience segmentation and like prediction KNWF files", "Confirm credentials are removed and nodes execute without cached-only results"),
            ("Dashboard", "PBIX, full PBIX, and PBIP project", "Refresh against the intended MySQL instance and replace instructional text boxes"),
            ("Report", "This editable DOCX", "Add student names, IDs, signatures, and export the approved copy to PDF"),
            ("README and supporting statements", "Repository files exist", "Populate setup, contribution, data-source, and AI-use text before submission"),
        ],
        widths=[Cm(3.2), Cm(6.0), Cm(7.4)],
        font_size=8.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    # Footers on all sections
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.clear()
        add_page_number(p)

    # Core properties
    doc.core_properties.title = "Segment Based Movie Promotion and Content Discovery Strategy"
    doc.core_properties.subject = "MovieLens 10M final project report"
    doc.core_properties.author = "MovieLens Content Strategy Project Team"
    doc.core_properties.keywords = "MovieLens, MySQL, KNIME, Power BI, K Means, Logistic Regression"
    doc.core_properties.comments = "Final project report generated from verified repository evidence"

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    figures = make_figures()
    build_report(figures)
