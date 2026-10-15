"""Shared, timeout-safe MySQL helpers for the MovieLens pipeline."""

from __future__ import annotations

import os
import re
import shutil
import sys
import threading
import time
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

import mysql.connector
from mysql.connector import Error
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
DEFAULT_DATA_DIR = REPO_ROOT / "data" / "ml-10M100K"
DEFAULT_RESULTS_DIR = SCRIPT_DIR / "results"
_IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def load_dotenv(path: Path | None = None) -> None:
    """Load a small .env file without adding a python-dotenv dependency."""
    env_path = path or SCRIPT_DIR / ".env"
    if not env_path.exists():
        return
    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip("\"").strip("'")
        os.environ.setdefault(key, value)


load_dotenv()


def database_name() -> str:
    return os.getenv("MYSQL_DATABASE", "movielens_strategy")


def data_dir() -> Path:
    configured = os.getenv("MOVIELENS_DATA_DIR")
    return Path(configured).expanduser().resolve() if configured else DEFAULT_DATA_DIR


def results_dir() -> Path:
    configured = os.getenv("MOVIELENS_RESULTS_DIR")
    return Path(configured).expanduser().resolve() if configured else DEFAULT_RESULTS_DIR


def ident(value: str) -> str:
    if not _IDENTIFIER.fullmatch(value):
        raise ValueError(f"Unsafe MySQL identifier: {value!r}")
    return f"`{value}`"


def connection_config(*, with_database: bool = True) -> dict[str, Any]:
    config: dict[str, Any] = {
        "host": os.getenv("MYSQL_HOST", "127.0.0.1"),
        "port": int(os.getenv("MYSQL_PORT", "3306")),
        "user": os.getenv("MYSQL_USER", "root"),
        "password": os.getenv("MYSQL_PASSWORD", ""),
        "charset": "utf8mb4",
        "use_unicode": True,
        "autocommit": True,
        "allow_local_infile": True,
        "connection_timeout": 30,
        # None means wait indefinitely for an active query. This avoids the
        # 900-second client-side cutoff seen in MySQL Workbench.
        "read_timeout": None,
        "write_timeout": None,
    }
    if with_database:
        config["database"] = database_name()
    return config


def connect(*, with_database: bool = True):
    connection = mysql.connector.connect(**connection_config(with_database=with_database))
    cursor = connection.cursor()
    try:
        cursor.execute("SET SESSION innodb_lock_wait_timeout = 3600")
        cursor.execute("SET SESSION lock_wait_timeout = 3600")
    finally:
        cursor.close()
    return connection


@contextmanager
def db(*, with_database: bool = True) -> Iterator[Any]:
    connection = connect(with_database=with_database)
    try:
        yield connection
    finally:
        if connection.is_connected():
            connection.close()


def execute(connection: Any, sql: str, params: Sequence[Any] | None = None) -> None:
    cursor = connection.cursor()
    try:
        cursor.execute(sql, params or ())
    finally:
        cursor.close()


def executemany(connection: Any, sql: str, rows: Iterable[Sequence[Any]]) -> None:
    cursor = connection.cursor()
    try:
        cursor.executemany(sql, rows)
    finally:
        cursor.close()


def query_all(
    connection: Any, sql: str, params: Sequence[Any] | None = None
) -> tuple[list[str], list[tuple[Any, ...]]]:
    cursor = connection.cursor()
    try:
        cursor.execute(sql, params or ())
        columns = list(cursor.column_names)
        return columns, list(cursor.fetchall())
    finally:
        cursor.close()


def scalar(connection: Any, sql: str, params: Sequence[Any] | None = None) -> Any:
    _, rows = query_all(connection, sql, params)
    return rows[0][0] if rows else None


def table_exists(connection: Any, table: str) -> bool:
    return bool(
        scalar(
            connection,
            """
            SELECT COUNT(*)
            FROM information_schema.tables
            WHERE table_schema = %s AND table_name = %s
            """,
            (database_name(), table),
        )
    )


def index_exists(connection: Any, table: str, index_name: str) -> bool:
    return bool(
        scalar(
            connection,
            """
            SELECT COUNT(*)
            FROM information_schema.statistics
            WHERE table_schema = %s AND table_name = %s AND index_name = %s
            """,
            (database_name(), table, index_name),
        )
    )


def equivalent_index(connection: Any, table: str, columns: Sequence[str]) -> str | None:
    """Return an existing index with the same leading columns, regardless of name."""
    _, rows = query_all(
        connection,
        """
        SELECT index_name, seq_in_index, column_name
        FROM information_schema.statistics
        WHERE table_schema = %s AND table_name = %s
        ORDER BY index_name, seq_in_index
        """,
        (database_name(), table),
    )
    grouped: dict[str, list[str]] = {}
    for name, _, column in rows:
        grouped.setdefault(str(name), []).append(str(column))
    requested = list(columns)
    for name, indexed_columns in grouped.items():
        if indexed_columns[: len(requested)] == requested:
            return name
    return None


def ensure_index(table: str, index_name: str, columns: Sequence[str]) -> None:
    """Create one index and recover cleanly if the original socket is dropped."""
    with db() as connection:
        existing = equivalent_index(connection, table, columns)
        if existing:
            print(f"[skip] equivalent index {existing} already covers {index_name}")
            return
        column_sql = ", ".join(ident(column) for column in columns)
        sql = (
            f"ALTER TABLE {ident(table)} ADD INDEX {ident(index_name)} ({column_sql}), "
            "ALGORITHM=INPLACE, LOCK=NONE"
        )
        try:
            run_timed(
                f"create index {index_name}",
                lambda: execute(connection, sql),
                connection=connection,
                report_failure=False,
            )
        except Error as exc:
            if exc.errno not in {2006, 2013, 2055}:
                raise
            print(f"[wait] connection dropped ({exc.errno}); checking server-side DDL status")
            wait_for_index(table, index_name, sql)
            print(f"[done] {index_name} completed on the server")


def wait_for_index(table: str, index_name: str, ddl_sql: str) -> None:
    """Poll after Error 2013: MySQL often continues ALTER TABLE server-side."""
    retried = False
    while True:
        time.sleep(10)
        try:
            with db() as probe:
                if index_exists(probe, table, index_name):
                    return
                active = scalar(
                    probe,
                    """
                    SELECT COUNT(*) FROM information_schema.processlist
                    WHERE DB = %s AND INFO LIKE %s
                    """,
                    (database_name(), f"%{index_name}%"),
                )
                if active:
                    print(f"[wait] {index_name} is still building on MySQL...")
                    continue
                if retried:
                    raise RuntimeError(f"Index {index_name} was not created after retry")
                print(f"[retry] server is no longer building {index_name}; retrying once")
                retried = True
                execute(probe, ddl_sql)
                if index_exists(probe, table, index_name):
                    return
        except Error as exc:
            if exc.errno not in {2003, 2006, 2013, 2055}:
                raise
            print(f"[wait] MySQL not ready while checking {index_name}: {exc}")


def print_rows(title: str, columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    print(f"\n{title}")
    if not rows:
        print("(no rows)")
        return
    widths = [len(str(column)) for column in columns]
    for row in rows:
        for i, value in enumerate(row):
            widths[i] = min(60, max(widths[i], len(str(value))))
    print(" | ".join(str(c).ljust(widths[i]) for i, c in enumerate(columns)))
    print("-+-".join("-" * width for width in widths))
    for row in rows:
        print(" | ".join(str(v).ljust(widths[i])[: widths[i]] for i, v in enumerate(row)))


def _safe_excel_value(value: Any) -> Any:
    # Prevent database text from being interpreted as a workbook formula.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sheet_name(name: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", name).strip(" '") or "Result"
    base = base[:31]
    candidate = base
    number = 2
    while candidate.casefold() in used:
        suffix = f"_{number}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        number += 1
    used.add(candidate.casefold())
    return candidate


def write_results_xlsx(
    script_name: str,
    result_sets: Sequence[tuple[str, Sequence[str], Sequence[Sequence[Any]]]],
) -> Path:
    """Write all query result sets from one script to separate Excel sheets."""
    if not result_sets:
        raise ValueError("At least one result set is required")
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F1F1F")

    for result_number, (title, columns, rows) in enumerate(result_sets, start=1):
        sheet = workbook.create_sheet(_sheet_name(title, used_names))
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.append([str(column) for column in columns])
        for row in rows:
            sheet.append([_safe_excel_value(value) for value in row])

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 22

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center")

        for column_number, column_name in enumerate(columns, start=1):
            normalized = str(column_name).lower()
            letter = get_column_letter(column_number)
            values = [sheet.cell(row=i, column=column_number).value for i in range(1, sheet.max_row + 1)]
            max_length = min(60, max(10, *(len(str(value)) for value in values if value is not None)))
            sheet.column_dimensions[letter].width = max_length + 2
            if any(token in normalized for token in ("rate", "probability")):
                for cell in sheet[letter][1:]:
                    cell.number_format = "0.00%"
            elif "percent" in normalized:
                for cell in sheet[letter][1:]:
                    cell.number_format = "0.00"
            elif any(token in normalized for token in ("count", "rows", "support", "users", "movies", "ratings", "tags")):
                for cell in sheet[letter][1:]:
                    cell.number_format = "#,##0"
            elif any(token in normalized for token in ("average", "avg", "stddev")):
                for cell in sheet[letter][1:]:
                    cell.number_format = "0.0000"
            elif any(token in normalized for token in ("date", "time", "period", "timestamp")):
                for cell in sheet[letter][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"

        if rows and columns:
            reference = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
            table = Table(displayName=f"Result_{result_number:02d}", ref=reference)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
        elif columns:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    destination_dir = results_dir()
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / f"{Path(script_name).stem}_results.xlsx"
    temporary = destination.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(destination)
    print(f"\n[export] Excel results: {destination}")
    return destination


def require_tables(connection: Any, names: Sequence[str]) -> None:
    missing = [name for name in names if not table_exists(connection, name)]
    if missing:
        raise RuntimeError(f"Missing required tables: {', '.join(missing)}")


def swap_staging_table(connection: Any, target: str, staging: str) -> str | None:
    """Atomically publish staging and retain the old derived table for rollback."""
    backups = swap_staging_tables(connection, ((target, staging),))
    return backups.get(target)


def swap_staging_tables(
    connection: Any, pairs: Sequence[tuple[str, str]]
) -> dict[str, str]:
    """Publish several related staging tables in one atomic RENAME statement."""
    suffix = f"{time.strftime('%Y%m%d_%H%M%S')}_{time.time_ns() % 1_000_000:06d}"
    clauses: list[str] = []
    backups: dict[str, str] = {}
    for target, staging in pairs:
        if table_exists(connection, target):
            backup = f"{target}_backup_{suffix}"
            clauses.append(f"{ident(target)} TO {ident(backup)}")
            backups[target] = backup
        clauses.append(f"{ident(staging)} TO {ident(target)}")
    execute(connection, "RENAME TABLE " + ", ".join(clauses))
    for target, backup in backups.items():
        print(f"[info] rollback copy for {target} retained as {backup}")
    return backups


def _format_duration(seconds: float) -> str:
    total = int(seconds)
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _mysql_progress(connection_id: int | None, stop: threading.Event, state: dict[str, Any]) -> None:
    """Poll server state from a separate connection without blocking the query."""
    if connection_id is None:
        return
    probe = None
    try:
        config = connection_config()
        config.update(connection_timeout=3, read_timeout=3, write_timeout=3)
        probe = mysql.connector.connect(**config)
        cursor = probe.cursor()
        while not stop.wait(2):
            cursor.execute(
                "SELECT TIME, COALESCE(STATE, COMMAND) FROM information_schema.PROCESSLIST WHERE ID=%s",
                (connection_id,),
            )
            process = cursor.fetchone()
            if process:
                state["mysql_time"] = int(process[0] or 0)
                state["mysql_state"] = str(process[1] or "running")
            try:
                cursor.execute(
                    """
                    SELECT es.EVENT_NAME, es.WORK_COMPLETED, es.WORK_ESTIMATED
                    FROM performance_schema.events_stages_current es
                    JOIN performance_schema.threads t ON t.THREAD_ID=es.THREAD_ID
                    WHERE t.PROCESSLIST_ID=%s
                    """,
                    (connection_id,),
                )
                stage = cursor.fetchone()
                if stage:
                    state["mysql_state"] = str(stage[0] or state.get("mysql_state", "running"))
                    completed, estimated = stage[1], stage[2]
                    if estimated and int(estimated) > 0:
                        state["percent"] = min(100.0, 100.0 * int(completed or 0) / int(estimated))
            except Error:
                # performance_schema stage consumers may be disabled or restricted.
                pass
    except Error:
        state["mysql_state"] = "running (server progress unavailable)"
    finally:
        if probe is not None and probe.is_connected():
            probe.close()


def _progress_display(
    label: str,
    started: float,
    stop: threading.Event,
    state: dict[str, Any],
) -> None:
    interactive = sys.stdout.isatty()
    interval = 0.25 if interactive else 15.0
    width = 24
    tick = 0
    while not stop.wait(interval):
        elapsed = time.monotonic() - started
        percent = state.get("percent")
        if percent is not None:
            filled = min(width, int(width * float(percent) / 100))
            bar = "#" * filled + "-" * (width - filled)
            progress = f"[{bar}] {float(percent):5.1f}%"
        else:
            cycle = 2 * width - 2
            position = tick % cycle
            if position >= width:
                position = cycle - position
            chars = ["-"] * width
            chars[position] = ">"
            progress = f"[{''.join(chars)}] running"
        mysql_state = str(state.get("mysql_state", "waiting for MySQL"))
        message = f"[work] {label} {progress} {_format_duration(elapsed)} | {mysql_state}"
        if interactive:
            terminal_width = shutil.get_terminal_size(fallback=(120, 24)).columns
            sys.stdout.write("\r" + message[: max(20, terminal_width - 1)].ljust(max(20, terminal_width - 1)))
            sys.stdout.flush()
        else:
            print(message, flush=True)
        tick += 1


def run_timed(
    label: str,
    action: Any,
    *,
    connection: Any | None = None,
    report_failure: bool = True,
) -> Any:
    print(f"\n[run ] {label}")
    started = time.monotonic()
    stop = threading.Event()
    state: dict[str, Any] = {}
    connection_id = None
    if connection is not None:
        try:
            connection_id = int(connection.connection_id)
        except (AttributeError, TypeError, ValueError):
            pass
    poller = threading.Thread(
        target=_mysql_progress,
        args=(connection_id, stop, state),
        daemon=True,
    )
    display = threading.Thread(
        target=_progress_display,
        args=(label, started, stop, state),
        daemon=True,
    )
    poller.start()
    display.start()
    try:
        result = action()
    except BaseException:
        elapsed = time.monotonic() - started
        stop.set()
        poller.join(timeout=0.5)
        display.join(timeout=0.5)
        if sys.stdout.isatty():
            sys.stdout.write("\r" + " " * max(20, shutil.get_terminal_size(fallback=(120, 24)).columns - 1) + "\r")
        if report_failure:
            print(f"[fail] {label} ({_format_duration(elapsed)})")
        raise
    stop.set()
    poller.join(timeout=0.5)
    display.join(timeout=0.5)
    if sys.stdout.isatty():
        sys.stdout.write("\r" + " " * max(20, shutil.get_terminal_size(fallback=(120, 24)).columns - 1) + "\r")
    print(f"[done] {label} ({_format_duration(time.monotonic() - started)})")
    return result

